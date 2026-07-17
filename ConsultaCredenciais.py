# === IMPORTACOES ===
import json
import getpass
import logging
import math
import os
import platform
import re
import socket
import subprocess
import sys
import tempfile
import time
import unicodedata
import urllib3
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from urllib.parse import urlparse

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from selenium import webdriver
from selenium.common.exceptions import (
    NoSuchElementException,
    TimeoutException,
    WebDriverException,
)
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

script_dir = os.path.dirname(sys.executable) if getattr(sys, "frozen", False) else os.path.dirname(os.path.abspath(__file__))
arquivo_excel = os.path.join(script_dir, "Credenciais.xlsx")

LOGIN_URL_MS = "http://ms.exemplo.com.br/login"
LOGIN_URL_AD = "http://www.rh.exemplo.com.br"
SUCCESS_URL_AD = "https://exemplo.com.br/Portal"

COL_EMAIL_INFORMADO = "A"
COL_SENHA = "B"
COL_DOCUMENTO_INFORMADO = "C"
COL_CPF_RESOLVIDO = "D"
COL_NOME = "E"
COL_EMAIL_FUNCIONAL = "F"
COL_RESULTADO_MS = "G"
COL_MSG_MS = "H"
COL_RESULTADO_AD = "I"
COL_MSG_AD = "J"
COL_ORIGEM_CPF = "K"
COL_STATUS_PROCESSAMENTO = "L"
COL_DATA_TESTE = "M"
COL_URL_VAZAMENTO = "N"

HEADER_JA_IDENTIFICADO_MES_ANTIGO = "ja identificado neste mes"
HEADER_CONTAGEM_IDENTIFICACOES = "Identificado quantas vezes"

HEADERS = [
    "E-mail informado",
    "Senha de entrada",
    "CPF ou RE informado",
    "CPF resolvido",
    "Nome completo",
    "E-mail funcional",
    "Resultado MS",
    "Mensagem MS",
    "Resultado AD",
    "Mensagem AD",
    "Origem do CPF",
    "Status do processamento",
    "Data e hora do teste",
    "URL de vazamento",
]

HEADERS_SAIDA = [
    "Nome",
    "E-mail",
    "CPF",
    "Senha",
    "MS",
    "AD",
    "Mensagem do MS",
    "Mensagem do AD",
    "URL de vazamento",
]

HEADERS_MENSAL = [
    "Nome",
    "E-mail",
    "CPF",
    "Senha",
    "MS",
    "AD",
    "Mensagem do MS",
    "Mensagem do AD",
    HEADER_CONTAGEM_IDENTIFICACOES,
    "Primeira identificação",
    "Data desta identificação",
    "Dia da semana",
    "URL de vazamento",
]

HEADERS_MENSAL_CREDENCIAIS = [
    "NOME",
    "CPF",
    "SENHA",
    "EMAIL",
    "URL",
    "DATA COLETA",
    "QUANTIDADE DE IDENTIFICAÇÕES",
    "DATA DAS IDENTIFICAÇÕES",
    "FONTE",
    "ACESSO AD",
    "ACESSO MS",
    "SITUAÇÃO LEGAL",
    "OBSERVAÇÕES",
    "RDS",
    "MSG BLOQUEIO",
]

HEADERS_MENSAL_TOTAL = [
    "NOME",
    "CPF",
    "SENHA",
    "EMAIL",
    "URL",
    "DATA COLETA",
    "QUANTIDADE DE IDENTIFICAÇÕES",
    "DATA DAS IDENTIFICAÇÕES",
    "FONTE",
    "ACESSO AD",
    "ACESSO MS",
    "SITUAÇÃO LEGAL",
]

ABAS_MENSAIS = {
    "AD": "Credenciais AD",
    "MS": "Credenciais MS",
    "Total": "Total",
}
MESES_PT = {
    1: "JAN",
    2: "FEV",
    3: "MAR",
    4: "ABR",
    5: "MAI",
    6: "JUN",
    7: "JUL",
    8: "AGO",
    9: "SET",
    10: "OUT",
    11: "NOV",
    12: "DEZ",
}

DIAS_SEMANA_PT = {
    0: "Segunda-feira",
    1: "Terça-feira",
    2: "Quarta-feira",
    3: "Quinta-feira",
    4: "Sexta-feira",
    5: "Sábado",
    6: "Domingo",
}

RESULTADO_SIM = "SIM"
RESULTADO_NAO = "NÃO"
RESULTADO_ERRO = "ERRO"
RESULTADO_NT = "N/T"
RESULTADO_INCONCLUSIVO = "INCONCLUSIVO"

SUCESSO = "SUCESSO"
CREDENCIAL_INVALIDA = "CREDENCIAL_INVALIDA"
CONTA_BLOQUEADA = "CONTA_BLOQUEADA"
CONTA_DESATIVADA = "CONTA_DESATIVADA"
SENHA_EXPIRADA = "SENHA_EXPIRADA"
ALTERACAO_DE_SENHA_OBRIGATORIA = "ALTERACAO_DE_SENHA_OBRIGATORIA"
ACESSO_NEGADO = "ACESSO_NEGADO"
USUARIO_NAO_ENCONTRADO = "USUARIO_NAO_ENCONTRADO"
SISTEMA_INDISPONIVEL = "SISTEMA_INDISPONIVEL"
ERRO_DE_REDE = "ERRO_DE_REDE"
TIMEOUT = "TIMEOUT"
ELEMENTO_NAO_ENCONTRADO = "ELEMENTO_NAO_ENCONTRADO"
PAGINA_INESPERADA = "PAGINA_INESPERADA"
ERRO_SELENIUM = "ERRO_SELENIUM"
ERRO_INTERNO = "ERRO_INTERNO"
DADOS_AUSENTES = "DADOS_AUSENTES"
NAO_TESTADO = "NAO_TESTADO"
INCONCLUSIVO = "INCONCLUSIVO"

CPF_INFORMADO = "CPF_INFORMADO"
CPF_RECONSTRUIDO_ZERO_INICIAL = "CPF_RECONSTRUIDO_ZERO_INICIAL"
RE_API = "RE_API"
RE_RECONSTRUIDO_ZERO_INICIAL = "RE_RECONSTRUIDO_ZERO_INICIAL"
EMAIL_DSQUERY = "EMAIL_DSQUERY"
NAO_ENCONTRADO = "NAO_ENCONTRADO"
AMBIGUO = "AMBIGUO"

CONCLUIDO = "CONCLUIDO"
CONCLUIDO_COM_RESSALVA = "CONCLUIDO_COM_RESSALVA"
ERRO_DE_ENTRADA = "ERRO_DE_ENTRADA"
ERRO_TECNICO = "ERRO_TECNICO"

CLASSIFICACOES_FUNCIONAIS = {
    CREDENCIAL_INVALIDA,
    CONTA_BLOQUEADA,
    CONTA_DESATIVADA,
    SENHA_EXPIRADA,
    ALTERACAO_DE_SENHA_OBRIGATORIA,
    ACESSO_NEGADO,
    USUARIO_NAO_ENCONTRADO,
}

CLASSIFICACOES_TECNICAS = {
    SISTEMA_INDISPONIVEL,
    ERRO_DE_REDE,
    TIMEOUT,
    ELEMENTO_NAO_ENCONTRADO,
    PAGINA_INESPERADA,
    ERRO_SELENIUM,
    ERRO_INTERNO,
}


class SeletoresMS:
    CPF = [(By.NAME, "vUSRNUMCPFAUX")]
    SENHA = [(By.NAME, "vSENHA")]
    BOTAO_LOGIN = [(By.NAME, "BTN_LOGIN")]
    MENSAGEM_ERRO = [(By.CSS_SELECTOR, "span#gxErrorViewer > div")]


class SeletoresAD:
    DROPDOWN_LOGIN = [(By.CLASS_NAME, "dropdown-toggle")]
    CPF = [(By.NAME, "CPF")]
    SENHA = [(By.CLASS_NAME, "password")]
    BOTAO_AUTENTICAR = [(By.ID, "btnAutenticar")]
    MENSAGEM_AUTENTICACAO = [(By.ID, "msgAutenticacao")]


@dataclass
class ResultadoOperacao:
    status: str
    mensagem: str
    valor: object = None
    detalhe_tecnico: str = ""


@dataclass
class ResultadoAuth:
    classificacao: str
    resultado_planilha: str
    mensagem: str
    detalhe_tecnico: str = ""


@dataclass
class IdentificadorNormalizado:
    tipo: str
    valor: str
    origem: str
    mensagem: str = ""


@dataclass
class ResultadoComando:
    args: list
    returncode: int
    stdout: str = ""
    stderr: str = ""
    timeout: bool = False


@dataclass
class UsuarioAD:
    dn: str
    atributos: dict

    @property
    def email(self):
        return normalizar_email(self.atributos.get("mail")) or normalizar_email(self.atributos.get("userPrincipalName"))

    @property
    def re(self):
        for chave in ("employeeNumber", "employeeID", "extensionAttribute1", "extensionAttribute2"):
            resultado = normalizar_re(self.atributos.get(chave))
            if resultado.status == SUCESSO:
                return resultado.valor
        return None

    @property
    def cpf(self):
        for valor in self.atributos.values():
            resultado = normalizar_cpf(valor)
            if resultado.status == SUCESSO:
                return resultado.valor
        return None

    @property
    def nome(self):
        return self.atributos.get("displayName") or self.atributos.get("cn") or self.atributos.get("name")

    @property
    def login(self):
        return self.atributos.get("sAMAccountName") or self.atributos.get("userPrincipalName")


@dataclass
class EventoAuditoria:
    categoria: str
    timestamp: datetime
    linha: object
    sistema: str
    cpf: str
    operacao: str
    status: str
    mensagem: str
    email: str = ""
    re_militar: str = ""


@dataclass
class ContextoLogs:
    execution_id: str
    inicio: datetime
    maquina: str
    ip: str
    usuario: str
    eventos: list
    falhas_log: list


LOG_CONTEXT = None


class ContextFilter(logging.Filter):
    def __init__(self, execution_id):
        super().__init__()
        self.execution_id = execution_id

    def filter(self, record):
        defaults = {
            "execution_id": self.execution_id,
            "fase": "-",
            "linha": "-",
            "sistema": "-",
            "cpf": "-",
            "operacao": "-",
            "status": "-",
            "duracao": "-",
        }
        for key, value in defaults.items():
            if not hasattr(record, key):
                setattr(record, key, value)
        return True


def configurar_logger(nome, caminho, execution_id):
    logger = logging.getLogger(nome)
    logger.setLevel(logging.INFO)
    for handler in logger.handlers:
        handler.close()
    logger.handlers.clear()
    handler = logging.NullHandler()
    handler.addFilter(ContextFilter(execution_id))
    logger.addHandler(handler)
    logger.propagate = False
    return logger


def obter_nome_maquina():
    try:
        return socket.gethostname() or platform.node() or "não identificado"
    except Exception:
        return "não identificado"


def obter_usuario_logado():
    try:
        usuario = os.environ.get("USERNAME") or getpass.getuser()
        dominio = os.environ.get("USERDOMAIN")
        if usuario and dominio:
            return f"{dominio}\\{usuario}"
        return usuario or "não identificado"
    except Exception:
        return "não identificado"


def obter_ip_local():
    candidatos = []
    try:
        hostname = socket.gethostname()
        for info in socket.getaddrinfo(hostname, None, socket.AF_INET):
            ip = info[4][0]
            if ip and not ip.startswith("127.") and ip not in candidatos:
                candidatos.append(ip)
    except Exception:
        pass
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            sock.connect(("10.255.255.255", 1))
            ip = sock.getsockname()[0]
            if ip and not ip.startswith("127.") and ip not in candidatos:
                candidatos.insert(0, ip)
        finally:
            sock.close()
    except Exception:
        pass
    return candidatos[0] if candidatos else "não identificado"


def iniciar_contexto_logs(execution_id, inicio):
    global LOG_CONTEXT
    LOG_CONTEXT = ContextoLogs(
        execution_id=execution_id,
        inicio=inicio,
        maquina=obter_nome_maquina(),
        ip=obter_ip_local(),
        usuario=obter_usuario_logado(),
        eventos=[],
        falhas_log=[],
    )
    return LOG_CONTEXT


def categoria_auditoria(fase, sistema, operacao, status):
    fase = str(fase or "").upper()
    operacao = str(operacao or "").lower()
    sistema = str(sistema or "").upper()
    if fase in ("IDENTIDADE", "API") or sistema in ("API", "AD", "DSQUERY") and "login" not in operacao:
        return "identidade"
    if fase == "AUTENTICACAO" or sistema in ("MS", "AD", "SELENIUM") and ("login" in operacao or "driver" in operacao or "auth" in operacao):
        return "credenciais"
    if fase in ("PLANILHA", "MENSAL") or "mensal" in operacao or "planilha" in operacao:
        return "planilha"
    return "outros"


def deve_auditar_evento(level, fase, sistema, operacao, status):
    if level >= logging.WARNING:
        return True
    status = str(status or "")
    if status in CLASSIFICACOES_TECNICAS or status in (INCONCLUSIVO, AMBIGUO, ERRO_DE_ENTRADA, ERRO_TECNICO):
        return True
    if str(fase or "").upper() == "AUTENTICACAO" and status in (TIMEOUT, ERRO_SELENIUM, ERRO_INTERNO, PAGINA_INESPERADA, ELEMENTO_NAO_ENCONTRADO):
        return True
    return False


def higienizar_mensagem_log(mensagem):
    texto = str(mensagem or "")
    texto = re.sub(r"(?i)(senha|password)\s*[:=]\s*\S+", r"\1=<removido>", texto)
    texto = re.sub(r"(?i)(token|cookie|authorization|bearer)\s*[:=]\s*[^\s,;]+", r"\1=<removido>", texto)
    texto = re.sub(r"(?i)([?&](?:token|senha|password|cookie|auth|authorization)=)[^&\s]+", r"\1<removido>", texto)
    return texto


def registrar_evento_auditoria(categoria, linha="-", sistema="-", cpf=None, operacao="-", status="-", mensagem="", email="", re_militar=""):
    if not LOG_CONTEXT:
        return
    LOG_CONTEXT.eventos.append(EventoAuditoria(
        categoria=categoria,
        timestamp=datetime.now(),
        linha=linha if linha is not None else "-",
        sistema=sistema or "-",
        cpf=mascarar_cpf(cpf),
        operacao=operacao or "-",
        status=status or "-",
        mensagem=higienizar_mensagem_log(mensagem),
        email=mascarar_email(email) if email else "",
        re_militar=mascarar_re(re_militar) if re_militar else "",
    ))


def log_event(logger, level, fase, linha, sistema, cpf, operacao, status, duracao, mensagem, exc_info=False):
    extra = {
        "fase": fase,
        "linha": linha if linha is not None else "-",
        "sistema": sistema or "-",
        "cpf": mascarar_cpf(cpf),
        "operacao": operacao,
        "status": status,
        "duracao": f"{duracao:.3f}s" if isinstance(duracao, (int, float)) else "-",
    }
    logger.log(level, mensagem, extra=extra, exc_info=exc_info)
    if deve_auditar_evento(level, fase, sistema, operacao, status):
        categoria = categoria_auditoria(fase, sistema, operacao, status)
        registrar_evento_auditoria(categoria, linha, sistema, cpf, operacao, status, mensagem)


def adicionar_erro(resumo, chave, linha, mensagem):
    resumo.setdefault(chave, [])
    resumo[chave].append((linha, mensagem))
    categoria = "identidade" if "identidade" in chave else "credenciais" if "autenticacao" in chave else "outros"
    registrar_evento_auditoria(categoria, linha=linha, operacao=chave, status="ERRO", mensagem=mensagem)


def texto_sem_acentos(valor):
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(ch for ch in texto if not unicodedata.combining(ch))


def normalizar_header(valor):
    return re.sub(r"[^a-z0-9]+", "", texto_sem_acentos(valor).lower())


def mascarar_cpf(cpf):
    digitos = re.sub(r"\D", "", str(cpf or ""))
    if len(digitos) != 11:
        return "-"
    return f"***.***.***-{digitos[9:]}"


def normalizar_email(valor):
    if valor is None:
        return None
    email = str(valor).strip().lower()
    if not email:
        return None
    if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        return email
    return None


def mascarar_email(email):
    email = normalizar_email(email)
    if not email:
        return "-"
    usuario, dominio = email.split("@", 1)
    if len(usuario) <= 2:
        usuario_mascarado = usuario[0] + "*" if usuario else "*"
    else:
        usuario_mascarado = f"{usuario[0]}***{usuario[-1]}"
    return f"{usuario_mascarado}@{dominio}"


def mascarar_re(re_militar):
    resultado = normalizar_re(re_militar)
    if resultado.status != SUCESSO:
        return "-"
    return f"***{resultado.valor[-3:]}"


def limpar_texto_identificador(valor):
    if valor is None:
        return ""
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if math.isnan(valor) or math.isinf(valor):
            return ""
        if valor.is_integer() and abs(valor) < 10**15:
            return str(int(valor))
        return ""
    texto = str(valor).strip()
    texto = texto.replace("\u200b", "").replace("\ufeff", "").strip()
    if texto.startswith("'"):
        texto = texto[1:].strip()
    if texto.startswith("="):
        return ""
    if re.fullmatch(r"[+-]?\d+(\.0+)?", texto):
        return texto.split(".")[0].lstrip("+")
    if re.fullmatch(r"[+-]?\d+(\.\d+)?[eE][+-]?\d+", texto):
        try:
            decimal = Decimal(texto)
            if decimal == decimal.to_integral_value():
                return str(int(decimal))
        except (InvalidOperation, ValueError):
            return ""
    return re.sub(r"\D", "", texto)


def cpf_valido(cpf):
    cpf = re.sub(r"\D", "", str(cpf or ""))
    if len(cpf) != 11:
        return False
    if cpf == cpf[0] * 11:
        return False
    soma = sum(int(cpf[i]) * (10 - i) for i in range(9))
    dig1 = (soma * 10) % 11
    dig1 = 0 if dig1 == 10 else dig1
    soma = sum(int(cpf[i]) * (11 - i) for i in range(10))
    dig2 = (soma * 10) % 11
    dig2 = 0 if dig2 == 10 else dig2
    return dig1 == int(cpf[9]) and dig2 == int(cpf[10])


def normalizar_cpf(valor):
    digitos = limpar_texto_identificador(valor)
    if not digitos:
        return ResultadoOperacao(DADOS_AUSENTES, "CPF vazio.")
    if len(digitos) == 11 and cpf_valido(digitos):
        return ResultadoOperacao(SUCESSO, "CPF valido.", digitos)
    if len(digitos) == 10:
        reconstruido = "0" + digitos
        if cpf_valido(reconstruido):
            return ResultadoOperacao(SUCESSO, "CPF reconstruido com zero inicial.", reconstruido, "zero_inicial")
    if len(digitos) == 11:
        return ResultadoOperacao(ERRO_DE_ENTRADA, "CPF com digitos verificadores invalidos.")
    return ResultadoOperacao(ERRO_DE_ENTRADA, "CPF com quantidade incorreta de digitos.")


def normalizar_re(valor):
    if valor is None:
        return ResultadoOperacao(DADOS_AUSENTES, "RE vazio.")
    if isinstance(valor, float):
        if math.isnan(valor) or math.isinf(valor):
            return ResultadoOperacao(DADOS_AUSENTES, "RE vazio.")
        if valor.is_integer() and abs(valor) < 10**9:
            valor = str(int(valor))
        else:
            return ResultadoOperacao(ERRO_DE_ENTRADA, "RE numerico invalido.")
    elif isinstance(valor, int):
        valor = str(valor)
    else:
        valor = str(valor).strip()

    if not valor:
        return ResultadoOperacao(DADOS_AUSENTES, "RE vazio.")
    if valor.startswith("'"):
        valor = valor[1:].strip()
    if valor.startswith("="):
        return ResultadoOperacao(ERRO_DE_ENTRADA, "RE com formato invalido.")
    if re.fullmatch(r"\d+\.0+", valor):
        valor = valor.split(".")[0]

    match = re.fullmatch(r"(?P<principal>\d{6})(?:-?\d)?|(?P<principal_pontuado>\d{3}\.\d{3})(?:-\d)?", valor)
    if not match:
        return ResultadoOperacao(ERRO_DE_ENTRADA, "RE com formato invalido.")
    re_normalizado = (match.group("principal") or match.group("principal_pontuado")).replace(".", "")
    if len(re_normalizado) != 6:
        return ResultadoOperacao(ERRO_DE_ENTRADA, "RE com quantidade incorreta de digitos.")
    return ResultadoOperacao(SUCESSO, "RE valido.", re_normalizado)


def normalizar_identificador(valor):
    digitos = limpar_texto_identificador(valor)
    if not digitos:
        return ResultadoOperacao(DADOS_AUSENTES, "CPF/RE ausente.")
    if len(digitos) in (10, 11):
        resultado_cpf = normalizar_cpf(valor)
        if resultado_cpf.status == SUCESSO:
            origem = CPF_RECONSTRUIDO_ZERO_INICIAL if resultado_cpf.detalhe_tecnico == "zero_inicial" else CPF_INFORMADO
            return ResultadoOperacao(SUCESSO, resultado_cpf.mensagem, IdentificadorNormalizado("CPF", resultado_cpf.valor, origem))
        return resultado_cpf
    resultado_re = normalizar_re(valor)
    if resultado_re.status == SUCESSO:
        return ResultadoOperacao(SUCESSO, resultado_re.mensagem, IdentificadorNormalizado("RE", resultado_re.valor, RE_API))
    if len(digitos) in (5, 6, 7):
        return resultado_re
    return ResultadoOperacao(ERRO_DE_ENTRADA, "Identificador nao e CPF nem RE reconhecido.")


class ClienteAPIsInternas:
    def __init__(self, logger):
        self.session = requests.Session()
        self.logger = logger
        self.timeout = (5, 15)
        self.max_tentativas = 3

    def fechar(self):
        self.session.close()

    def _get_json(self, url, operacao, linha=None, cpf=None):
        inicio = time.perf_counter()
        for tentativa in range(1, self.max_tentativas + 1):
            try:
                response = self.session.get(url, verify=False, timeout=self.timeout)
                if response.status_code in (429, 500, 502, 503, 504):
                    if tentativa < self.max_tentativas:
                        time.sleep(0.5 * tentativa)
                        continue
                    return ResultadoOperacao(SISTEMA_INDISPONIVEL, "API indisponivel.", detalhe_tecnico=f"HTTP {response.status_code}")
                if response.status_code in (401, 403):
                    return ResultadoOperacao(ACESSO_NEGADO, "Acesso negado pela API.", detalhe_tecnico=f"HTTP {response.status_code}")
                if response.status_code == 404:
                    return ResultadoOperacao(NAO_ENCONTRADO, "Registro nao encontrado pela API.", detalhe_tecnico="HTTP 404")
                response.raise_for_status()
                if not response.content:
                    return ResultadoOperacao(INCONCLUSIVO, "Resposta vazia da API.")
                try:
                    return ResultadoOperacao(SUCESSO, "Consulta realizada.", response.json())
                except ValueError as exc:
                    return ResultadoOperacao(ERRO_INTERNO, "JSON invalido retornado pela API.", detalhe_tecnico=str(exc))
            except requests.Timeout as exc:
                if tentativa < self.max_tentativas:
                    time.sleep(0.5 * tentativa)
                    continue
                return ResultadoOperacao(TIMEOUT, "Tempo limite excedido consultando API.", detalhe_tecnico=str(exc))
            except requests.ConnectionError as exc:
                if tentativa < self.max_tentativas:
                    time.sleep(0.5 * tentativa)
                    continue
                return ResultadoOperacao(ERRO_DE_REDE, "Erro de rede consultando API.", detalhe_tecnico=str(exc))
            except requests.RequestException as exc:
                return ResultadoOperacao(ERRO_DE_REDE, "Erro HTTP consultando API.", detalhe_tecnico=str(exc))
            finally:
                duracao = time.perf_counter() - inicio
                log_event(self.logger, logging.INFO, "API", linha, "API", cpf, operacao, "tentativa", duracao, f"Tentativa {tentativa}")
        return ResultadoOperacao(ERRO_INTERNO, "Erro inesperado na consulta da API.")

    def buscar_cpf_por_re(self, re_militar, linha=None):
        url = f"https://siteapi.com.br/api/v1//re/{re_militar}/dados"
        resposta = self._get_json(url, "buscar_cpf_por_re", linha=linha)
        if resposta.status != SUCESSO:
            return resposta
        dados = resposta.valor.get("dados") if isinstance(resposta.valor, dict) else None
        if not isinstance(dados, list) or not dados:
            return ResultadoOperacao(NAO_ENCONTRADO, "CPF nao encontrado pelo RE.")
        cpf = dados[0].get("cpf", {}).get("cpfComDigito")
        cpf_normalizado = normalizar_cpf(cpf)
        if cpf_normalizado.status != SUCESSO:
            return ResultadoOperacao(INCONCLUSIVO, "API retornou CPF invalido para o RE.", detalhe_tecnico=cpf_normalizado.mensagem)
        return ResultadoOperacao(SUCESSO, "CPF encontrado pelo RE.", cpf_normalizado.valor)

    def buscar_nome_por_cpf(self, cpf, linha=None):
        url = f"https://siteapi.com.br/api/v1/PolicialMilitar/cpf/{cpf}/dados"
        resposta = self._get_json(url, "buscar_nome_por_cpf", linha=linha, cpf=cpf)
        if resposta.status != SUCESSO:
            return resposta
        dados = resposta.valor.get("dados") if isinstance(resposta.valor, dict) else None
        if not isinstance(dados, list) or not dados:
            return ResultadoOperacao(NAO_ENCONTRADO, "Nome nao encontrado pelo CPF.")
        nome = dados[0].get("nomeCompleto")
        if not nome:
            return ResultadoOperacao(NAO_ENCONTRADO, "Nome ausente na resposta da API.")
        return ResultadoOperacao(SUCESSO, "Nome encontrado.", str(nome).strip())

    def buscar_email_por_cpf(self, cpf, linha=None):
        url = f"https://siteapi.com.br/api/v1/PolicialMilitar/cpf/{cpf}/informacaoContato"
        resposta = self._get_json(url, "buscar_email_por_cpf", linha=linha, cpf=cpf)
        if resposta.status != SUCESSO:
            return resposta
        dados = resposta.valor.get("dados") if isinstance(resposta.valor, dict) else None
        if not isinstance(dados, list) or not dados:
            return ResultadoOperacao(NAO_ENCONTRADO, "E-mail funcional nao encontrado pelo CPF.")
        for email_info in dados[0].get("emails", []):
            if email_info.get("tipoContato", {}).get("identificador") == 4:
                endereco = email_info.get("endereco")
                if endereco:
                    return ResultadoOperacao(SUCESSO, "E-mail funcional encontrado.", str(endereco).strip())
        return ResultadoOperacao(NAO_ENCONTRADO, "E-mail funcional nao encontrado pelo CPF.")


AD_ATTRS = [
    "mail",
    "userPrincipalName",
    "sAMAccountName",
    "employeeNumber",
    "employeeID",
    "displayName",
    "cn",
    "name",
    "description",
    "info",
    "extensionAttribute1",
    "extensionAttribute2",
    "extensionAttribute3",
    "extensionAttribute4",
    "extensionAttribute5",
    "extensionAttribute6",
    "extensionAttribute7",
    "extensionAttribute8",
    "extensionAttribute9",
    "extensionAttribute10",
    "extensionAttribute11",
    "extensionAttribute12",
    "extensionAttribute13",
    "extensionAttribute14",
    "extensionAttribute15",
    "cpf",
]


def escapar_filtro_ldap(valor):
    texto = str(valor)
    return (
        texto.replace("\\", r"\5c")
        .replace("*", r"\2a")
        .replace("(", r"\28")
        .replace(")", r"\29")
        .replace("\x00", r"\00")
    )


def executar_comando_externo(args, timeout=20):
    try:
        cmd = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout, shell=False)
        return ResultadoComando(args=args, returncode=cmd.returncode, stdout=cmd.stdout or "", stderr=cmd.stderr or "")
    except FileNotFoundError as exc:
        return ResultadoComando(args=args, returncode=-1, stderr=str(exc))
    except subprocess.TimeoutExpired as exc:
        stdout = exc.stdout.decode("utf-8", "replace") if isinstance(exc.stdout, bytes) else (exc.stdout or "")
        stderr = exc.stderr.decode("utf-8", "replace") if isinstance(exc.stderr, bytes) else (exc.stderr or "")
        return ResultadoComando(args=args, returncode=-1, stdout=stdout, stderr=stderr, timeout=True)


def classificar_erro_dsquery(resultado):
    texto = texto_sem_acentos(f"{resultado.stdout}\n{resultado.stderr}").lower()
    if resultado.timeout:
        return TIMEOUT, "Tempo limite excedido consultando Active Directory."
    if resultado.returncode == -1 and (
        "nao encontrado" in texto
        or "not found" in texto
        or "cannot find" in texto
        or "nao pode encontrar" in texto
    ):
        return ERRO_DE_REDE, "dsquery nao encontrado neste ambiente."
    if "access is denied" in texto or "acesso negado" in texto:
        return ACESSO_NEGADO, "Acesso negado ao consultar Active Directory."
    if "server is not operational" in texto or "domain" in texto or "dominio" in texto or "contatado" in texto or "specified domain" in texto:
        return ERRO_DE_REDE, "Falha de dominio ao consultar Active Directory."
    if resultado.returncode != 0:
        return ERRO_DE_REDE, "Erro ao consultar Active Directory."
    return INCONCLUSIVO, "Consulta ao Active Directory inconclusiva."


def log_ad(logger, linha, tipo_busca, valor_mascarado, status, mensagem, inicio, returncode=None, stderr=False, objetos=None):
    detalhe = f"busca={tipo_busca}; valor={valor_mascarado}; {mensagem}"
    if returncode is not None:
        detalhe += f"; returncode={returncode}"
    detalhe += f"; stderr={'sim' if stderr else 'nao'}"
    if objetos is not None:
        detalhe += f"; objetos={objetos}"
    log_event(logger, logging.INFO, "IDENTIDADE", linha, "AD", None, f"ad_{tipo_busca}", status, time.perf_counter() - inicio, detalhe)


def extrair_dns_dsquery(stdout):
    dns = []
    for linha in (stdout or "").splitlines():
        valor = linha.strip()
        if not valor:
            continue
        if valor.startswith('"') and valor.endswith('"'):
            valor = valor[1:-1]
        if valor.lower().startswith("cn=") and valor not in dns:
            dns.append(valor)
    return dns


def extrair_valor_atributo(stdout, atributo):
    linhas = [linha.rstrip() for linha in (stdout or "").splitlines() if linha.strip()]
    if len(linhas) < 2:
        return None
    if normalizar_header(linhas[0]) != normalizar_header(atributo):
        return None
    valor = linhas[1].strip()
    return None if valor in ("-", "") else valor


def buscar_dns_ad_por_filtro(filtro_ldap, logger, linha, tipo_busca, valor_mascarado, limite=20):
    inicio = time.perf_counter()
    args = ["dsquery", "*", "-filter", filtro_ldap, "-limit", str(limite)]
    resultado = executar_comando_externo(args, timeout=20)
    if resultado.returncode != 0 or resultado.timeout:
        status, mensagem = classificar_erro_dsquery(resultado)
        log_ad(logger, linha, tipo_busca, valor_mascarado, status, mensagem, inicio, resultado.returncode, bool(resultado.stderr), 0)
        return ResultadoOperacao(status, mensagem, detalhe_tecnico=resultado.stderr or resultado.stdout)
    dns = extrair_dns_dsquery(resultado.stdout)
    if resultado.stderr:
        log_ad(logger, linha, tipo_busca, valor_mascarado, INCONCLUSIVO, "stderr presente na busca de DN.", inicio, resultado.returncode, True, len(dns))
    else:
        log_ad(logger, linha, tipo_busca, valor_mascarado, SUCESSO if dns else NAO_ENCONTRADO, "DNs retornados pelo dsquery.", inicio, resultado.returncode, False, len(dns))
    if not dns:
        return ResultadoOperacao(NAO_ENCONTRADO, "Usuario nao encontrado no Active Directory.")
    return ResultadoOperacao(SUCESSO, "DNs encontrados no Active Directory.", dns)


def buscar_atributo_ad(dn, atributo, logger, linha):
    inicio = time.perf_counter()
    args = ["dsquery", "*", dn, "-scope", "base", "-attr", atributo]
    resultado = executar_comando_externo(args, timeout=20)
    if resultado.returncode != 0 or resultado.timeout:
        status, mensagem = classificar_erro_dsquery(resultado)
        log_ad(logger, linha, "atributo", atributo, status, mensagem, inicio, resultado.returncode, bool(resultado.stderr), 0)
        return ResultadoOperacao(status, mensagem, detalhe_tecnico=resultado.stderr or resultado.stdout)
    valor = extrair_valor_atributo(resultado.stdout, atributo)
    log_ad(logger, linha, "atributo", atributo, SUCESSO if valor else NAO_ENCONTRADO, f"atributo_localizado={bool(valor)}", inicio, resultado.returncode, bool(resultado.stderr), 1 if valor else 0)
    return ResultadoOperacao(SUCESSO, "Atributo consultado.", valor)


def carregar_usuario_ad_por_dn(dn, logger, linha):
    atributos = {}
    for atributo in AD_ATTRS:
        resposta = buscar_atributo_ad(dn, atributo, logger, linha)
        if resposta.status == SUCESSO and resposta.valor:
            atributos[atributo] = resposta.valor
    return UsuarioAD(dn=dn, atributos=atributos)


def carregar_usuarios_ad(dns, logger, linha):
    return [carregar_usuario_ad_por_dn(dn, logger, linha) for dn in dns]


def buscar_usuarios_ad_por_email(email, logger, linha=None):
    email_normalizado = normalizar_email(email)
    if not email_normalizado:
        return ResultadoOperacao(ERRO_DE_ENTRADA, "E-mail invalido para consulta ao Active Directory.")
    email_ldap = escapar_filtro_ldap(email_normalizado)
    filtro = f"(|(mail={email_ldap})(userPrincipalName={email_ldap}))"
    dns = buscar_dns_ad_por_filtro(filtro, logger, linha, "email", mascarar_email(email_normalizado))
    if dns.status != SUCESSO:
        return dns
    return ResultadoOperacao(SUCESSO, "Usuarios encontrados por e-mail.", carregar_usuarios_ad(dns.valor, logger, linha))


def buscar_usuarios_ad_por_re(re_militar, logger, linha=None):
    re_normalizado = normalizar_re(re_militar)
    if re_normalizado.status != SUCESSO:
        return re_normalizado
    re_ldap = escapar_filtro_ldap(re_normalizado.valor)
    filtro = f"(|(employeeNumber={re_ldap})(employeeID={re_ldap})(extensionAttribute1={re_ldap})(extensionAttribute2={re_ldap}))"
    dns = buscar_dns_ad_por_filtro(filtro, logger, linha, "re", mascarar_re(re_normalizado.valor))
    if dns.status != SUCESSO:
        return dns
    return ResultadoOperacao(SUCESSO, "Usuarios encontrados por RE.", carregar_usuarios_ad(dns.valor, logger, linha))


def buscar_usuarios_ad_por_cpf(cpf, logger, linha=None):
    cpf_normalizado = normalizar_cpf(cpf)
    if cpf_normalizado.status != SUCESSO:
        return cpf_normalizado
    cpf_ldap = escapar_filtro_ldap(cpf_normalizado.valor)
    filtro = (
        f"(|(employeeID={cpf_ldap})(cpf={cpf_ldap})(extensionAttribute3={cpf_ldap})(extensionAttribute4={cpf_ldap})"
        f"(description=*{cpf_ldap}*)(info=*{cpf_ldap}*))"
    )
    dns = buscar_dns_ad_por_filtro(filtro, logger, linha, "cpf", mascarar_cpf(cpf_normalizado.valor))
    if dns.status != SUCESSO:
        return dns
    return ResultadoOperacao(SUCESSO, "Usuarios encontrados por CPF.", carregar_usuarios_ad(dns.valor, logger, linha))


def deduplicar_usuarios_ad(usuarios):
    por_dn = {}
    for usuario in usuarios:
        if usuario.dn not in por_dn:
            por_dn[usuario.dn] = usuario
            continue
        por_dn[usuario.dn].atributos.update(usuario.atributos)
    return list(por_dn.values())


def correlacionar_usuarios_ad(usuarios, email=None, re_militar=None, cpf=None, logger=None, linha=None):
    usuarios = deduplicar_usuarios_ad(usuarios)
    if not usuarios:
        return ResultadoOperacao(NAO_ENCONTRADO, "Nenhum usuario localizado no Active Directory.")

    email_normalizado = normalizar_email(email)
    resultado_re = normalizar_re(re_militar) if re_militar else None
    resultado_cpf = normalizar_cpf(cpf) if cpf else None
    re_normalizado = resultado_re.valor if resultado_re and resultado_re.status == SUCESSO else None
    cpf_normalizado = resultado_cpf.valor if resultado_cpf and resultado_cpf.status == SUCESSO else None

    candidatos = usuarios
    if email_normalizado:
        correspondentes_email = [u for u in usuarios if u.email == email_normalizado]
        if correspondentes_email:
            candidatos = correspondentes_email
        elif len(usuarios) > 1:
            return ResultadoOperacao(AMBIGUO, "Multiplos usuarios retornados sem correspondencia exata de e-mail.", usuarios)

    if len(candidatos) > 1:
        return ResultadoOperacao(AMBIGUO, "Multiplos usuarios encontrados no Active Directory.", candidatos)

    usuario = candidatos[0]
    divergencias = []
    if email_normalizado and usuario.email and usuario.email != email_normalizado:
        divergencias.append(f"e-mail AD {mascarar_email(usuario.email)} diferente do informado {mascarar_email(email_normalizado)}")
    if re_normalizado and usuario.re and usuario.re != re_normalizado:
        divergencias.append(f"RE AD {mascarar_re(usuario.re)} diferente do informado {mascarar_re(re_normalizado)}")
    if cpf_normalizado and usuario.cpf and usuario.cpf != cpf_normalizado:
        divergencias.append(f"CPF AD {mascarar_cpf(usuario.cpf)} diferente do informado {mascarar_cpf(cpf_normalizado)}")
    if divergencias:
        mensagem = "Conflito de identidade no Active Directory: " + "; ".join(divergencias)
        if logger:
            log_event(logger, logging.WARNING, "IDENTIDADE", linha, "AD", cpf_normalizado, "correlacionar_ad", AMBIGUO, 0, mensagem)
        return ResultadoOperacao(AMBIGUO, mensagem, usuario)

    if not usuario.cpf and not usuario.re:
        return ResultadoOperacao(INCONCLUSIVO, "Usuario localizado no AD, mas sem CPF ou RE confiavel.", usuario)
    if logger:
        atributos = ", ".join(sorted(usuario.atributos.keys()))
        log_event(logger, logging.INFO, "IDENTIDADE", linha, "AD", usuario.cpf, "correlacionar_ad", SUCESSO, 0, f"Usuario AD aceito; atributos={atributos}; login={usuario.login or '-'}")
    return ResultadoOperacao(SUCESSO, "Usuario AD correlacionado com seguranca.", usuario)


def buscar_usuario_ad_com_identificadores(email=None, re_militar=None, cpf=None, logger=None, linha=None):
    usuarios = []
    erros_tecnicos = []
    if re_militar:
        resposta = buscar_usuarios_ad_por_re(re_militar, logger, linha)
        if resposta.status == SUCESSO:
            usuarios.extend(resposta.valor)
        elif resposta.status in CLASSIFICACOES_TECNICAS or resposta.status == ACESSO_NEGADO:
            erros_tecnicos.append(resposta.mensagem)
    if cpf:
        resposta = buscar_usuarios_ad_por_cpf(cpf, logger, linha)
        if resposta.status == SUCESSO:
            usuarios.extend(resposta.valor)
        elif resposta.status in CLASSIFICACOES_TECNICAS or resposta.status == ACESSO_NEGADO:
            erros_tecnicos.append(resposta.mensagem)
    if email:
        resposta = buscar_usuarios_ad_por_email(email, logger, linha)
        if resposta.status == SUCESSO:
            usuarios.extend(resposta.valor)
        elif resposta.status in CLASSIFICACOES_TECNICAS or resposta.status == ACESSO_NEGADO:
            erros_tecnicos.append(resposta.mensagem)

    correlacao = correlacionar_usuarios_ad(usuarios, email=email, re_militar=re_militar, cpf=cpf, logger=logger, linha=linha)
    if correlacao.status == NAO_ENCONTRADO and erros_tecnicos:
        return ResultadoOperacao(ERRO_DE_REDE, "Falha tecnica na consulta AD: " + "; ".join(erros_tecnicos))
    return correlacao


def buscar_re_por_email(email, logger, linha=None):
    resposta = buscar_usuario_ad_com_identificadores(email=email, logger=logger, linha=linha)
    if resposta.status != SUCESSO:
        return resposta
    usuario = resposta.valor
    if usuario.re:
        return ResultadoOperacao(SUCESSO, "RE encontrado por e-mail no Active Directory.", usuario.re)
    if usuario.cpf:
        return ResultadoOperacao(SUCESSO, "CPF encontrado por e-mail no Active Directory.", usuario.cpf, detalhe_tecnico="CPF")
    return ResultadoOperacao(INCONCLUSIVO, "Usuario encontrado no AD sem RE ou CPF.")


def validar_arquivo_entrada(caminho):
    if not os.path.exists(caminho):
        raise FileNotFoundError(f"Arquivo de entrada nao encontrado: {caminho}")
    if not caminho.lower().endswith(".xlsx"):
        raise ValueError("Arquivo de entrada precisa ter extensao .xlsx.")
    if os.path.getsize(caminho) == 0:
        raise ValueError("Arquivo de entrada esta vazio.")
    diretorio = os.path.dirname(caminho)
    if not os.access(caminho, os.R_OK):
        raise PermissionError("Sem permissao de leitura no arquivo de entrada.")
    if not os.access(diretorio, os.W_OK):
        raise PermissionError("Sem permissao de escrita no diretorio de saida.")
    teste_path = None
    try:
        fd, teste_path = tempfile.mkstemp(prefix=".credenciais_write_test_", suffix=".tmp", dir=diretorio)
        os.close(fd)
    finally:
        if teste_path and os.path.exists(teste_path):
            os.remove(teste_path)
    wb = load_workbook(caminho)
    ws = wb.active
    if ws.max_row < 2:
        raise ValueError("Planilha sem linhas de dados.")
    if ws.max_column < 3:
        raise ValueError("Planilha precisa ter ao menos as colunas A, B e C.")
    headers = [normalizar_header(ws.cell(1, col).value) for col in range(1, min(ws.max_column, len(HEADERS)) + 1)]
    layout_antigo = headers[:3] == ["email", "senha", "cpf"]
    headers_esperados = [normalizar_header(h) for h in HEADERS]
    layout_novo = headers[:len(HEADERS)] == headers_esperados or headers[:len(HEADERS) - 1] == headers_esperados[:-1]
    if not layout_antigo and not layout_novo:
        raise ValueError("Layout da planilha nao e compativel com o esperado.")
    return wb, ws, layout_novo


def localizar_coluna_url_vazamento(ws):
    header_url = normalizar_header("URL de vazamento")
    for col in range(1, ws.max_column + 1):
        if normalizar_header(ws.cell(1, col).value) == header_url:
            return col
    return None


def preparar_layout(ws, layout_novo):
    if layout_novo:
        if normalizar_header(ws[f"{COL_URL_VAZAMENTO}1"].value) != normalizar_header("URL de vazamento"):
            ws[f"{COL_URL_VAZAMENTO}1"] = "URL de vazamento"
        return
    coluna_url = localizar_coluna_url_vazamento(ws)
    linhas = []
    for row in range(2, ws.max_row + 1):
        linhas.append((
            ws[f"A{row}"].value,
            ws[f"B{row}"].value,
            ws[f"C{row}"].value,
            ws.cell(row, coluna_url).value if coluna_url else None,
        ))
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for offset, (email, senha, documento, url_vazamento) in enumerate(linhas, start=2):
        ws[f"{COL_EMAIL_INFORMADO}{offset}"] = email
        ws[f"{COL_SENHA}{offset}"] = senha
        ws[f"{COL_DOCUMENTO_INFORMADO}{offset}"] = documento
        ws[f"{COL_URL_VAZAMENTO}{offset}"] = url_vazamento
        for col in range(4, 14):
            ws.cell(offset, col).value = None


def aplicar_estilos(ws):
    success_fill = PatternFill("solid", fgColor="C6EFCE")
    failure_fill = PatternFill("solid", fgColor="FFC7CE")
    error_fill = PatternFill("solid", fgColor="FFEB9C")
    inconclusive_fill = PatternFill("solid", fgColor="D9EAF7")
    nt_fill = PatternFill("solid", fgColor="D9D9D9")
    fills = {
        RESULTADO_SIM: success_fill,
        RESULTADO_NAO: failure_fill,
        RESULTADO_ERRO: error_fill,
        RESULTADO_INCONCLUSIVO: inconclusive_fill,
        RESULTADO_NT: nt_fill,
    }
    for row in range(2, ws.max_row + 1):
        for col in [COL_RESULTADO_MS, COL_RESULTADO_AD]:
            cell = ws[f"{col}{row}"]
            if cell.value in fills:
                cell.fill = fills[cell.value]
            cell.alignment = Alignment(horizontal="center")


def aplicar_estilos_saida(ws):
    success_fill = PatternFill("solid", fgColor="C6EFCE")
    failure_fill = PatternFill("solid", fgColor="FFC7CE")
    for col in range(1, len(HEADERS_SAIDA) + 1):
        cell = ws.cell(1, col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in range(2, ws.max_row + 1):
        for col in (5, 6):
            cell = ws.cell(row, col)
            if cell.value == RESULTADO_SIM:
                cell.fill = success_fill
            elif cell.value == RESULTADO_NAO:
                cell.fill = failure_fill
            cell.alignment = Alignment(horizontal="center")
    larguras = {
        "A": 34,
        "B": 38,
        "C": 16,
        "D": 28,
        "E": 12,
        "F": 12,
        "G": 58,
        "H": 58,
        "I": 58,
    }
    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura


def criar_workbook_saida(ws_processado):
    wb_saida = Workbook()
    ws_saida = wb_saida.active
    ws_saida.title = "Resultado"
    for col, header in enumerate(HEADERS_SAIDA, start=1):
        ws_saida.cell(1, col).value = header
    for row in range(2, ws_processado.max_row + 1):
        linha_saida = row
        ws_saida.cell(linha_saida, 1).value = ws_processado[f"{COL_NOME}{row}"].value
        ws_saida.cell(linha_saida, 2).value = ws_processado[f"{COL_EMAIL_FUNCIONAL}{row}"].value
        ws_saida.cell(linha_saida, 3).value = ws_processado[f"{COL_CPF_RESOLVIDO}{row}"].value
        ws_saida.cell(linha_saida, 4).value = ws_processado[f"{COL_SENHA}{row}"].value
        ws_saida.cell(linha_saida, 5).value = ws_processado[f"{COL_RESULTADO_MS}{row}"].value
        ws_saida.cell(linha_saida, 6).value = ws_processado[f"{COL_RESULTADO_AD}{row}"].value
        ws_saida.cell(linha_saida, 7).value = ws_processado[f"{COL_MSG_MS}{row}"].value
        ws_saida.cell(linha_saida, 8).value = ws_processado[f"{COL_MSG_AD}{row}"].value
        ws_saida.cell(linha_saida, 9).value = ws_processado[f"{COL_URL_VAZAMENTO}{row}"].value
    aplicar_estilos_saida(ws_saida)
    return wb_saida


def resultado_planilha_por_classificacao(classificacao):
    if classificacao == SUCESSO:
        return RESULTADO_SIM
    if classificacao in CLASSIFICACOES_FUNCIONAIS:
        return RESULTADO_NAO
    if classificacao in CLASSIFICACOES_TECNICAS:
        return RESULTADO_ERRO
    if classificacao in (DADOS_AUSENTES, NAO_TESTADO):
        return RESULTADO_NT
    return RESULTADO_INCONCLUSIVO


def classificar_mensagem_autenticacao(mensagem):
    texto = texto_sem_acentos(mensagem).lower()
    if not texto:
        return INCONCLUSIVO
    if "senha invalida" in texto or "usuario ou senha incorreta" in texto or "usuario e senha incorreta" in texto:
        return CREDENCIAL_INVALIDA
    if "bloquead" in texto:
        return CONTA_BLOQUEADA
    if "inativo" in texto or "desativad" in texto:
        return CONTA_DESATIVADA
    if "senha expirada" in texto:
        return SENHA_EXPIRADA
    if "alteracao de senha" in texto or "troca de senha" in texto:
        return ALTERACAO_DE_SENHA_OBRIGATORIA
    if "nao cadastrado" in texto or "cpf invalido" in texto or "usuario nao encontrado" in texto:
        return USUARIO_NAO_ENCONTRADO
    if "acesso negado" in texto:
        return ACESSO_NEGADO
    return INCONCLUSIVO


def limpar_mensagem_autenticacao(mensagem):
    texto = "" if mensagem is None else str(mensagem)
    texto = re.sub(r"[\r\n\t]+", " ", texto)
    texto = re.sub(r"\s+", " ", texto).strip()
    texto = re.sub(r"\s*[×x]\s*$", "", texto, flags=re.IGNORECASE).strip()
    return texto


def mensagem_para_planilha(classificacao, mensagem_original=""):
    mensagem_original = limpar_mensagem_autenticacao(mensagem_original)
    if mensagem_original and classificacao in CLASSIFICACOES_FUNCIONAIS:
        return str(mensagem_original).strip()
    mensagens = {
        SUCESSO: "Autenticacao realizada com sucesso.",
        CREDENCIAL_INVALIDA: "Credencial invalida confirmada pela mensagem da pagina.",
        CONTA_BLOQUEADA: "Conta bloqueada.",
        CONTA_DESATIVADA: "Conta desativada ou inativa.",
        SENHA_EXPIRADA: "Senha expirada.",
        ALTERACAO_DE_SENHA_OBRIGATORIA: "Alteracao de senha obrigatoria.",
        ACESSO_NEGADO: "Acesso negado.",
        USUARIO_NAO_ENCONTRADO: "Usuario nao encontrado ou CPF invalido.",
        SISTEMA_INDISPONIVEL: "Sistema indisponivel.",
        ERRO_DE_REDE: "Erro de rede.",
        TIMEOUT: "Tempo limite excedido aguardando resposta.",
        ELEMENTO_NAO_ENCONTRADO: "Elemento esperado nao localizado.",
        PAGINA_INESPERADA: "Pagina retornada nao corresponde ao fluxo esperado.",
        ERRO_SELENIUM: "Erro tecnico do Selenium.",
        ERRO_INTERNO: "Erro interno durante o processamento.",
        DADOS_AUSENTES: "CPF ou senha ausente.",
        NAO_TESTADO: "Nao testado.",
        INCONCLUSIVO: "Nao foi possivel confirmar sucesso ou falha.",
    }
    return mensagens.get(classificacao, mensagens[INCONCLUSIVO])


def aguardar_primeiro_elemento_valido(driver, seletores, timeout=10, clicavel=False):
    ultimo_erro = None
    for seletor in seletores:
        try:
            condicao = EC.element_to_be_clickable(seletor) if clicavel else EC.presence_of_element_located(seletor)
            return WebDriverWait(driver, timeout).until(condicao)
        except TimeoutException as exc:
            ultimo_erro = exc
    raise TimeoutException(str(ultimo_erro) if ultimo_erro else "Elemento nao localizado.")


def elemento_existe(driver, seletores):
    for by, value in seletores:
        try:
            if driver.find_elements(by, value):
                return True
        except WebDriverException:
            return False
    return False


def texto_visivel_primeiro_elemento(driver, seletores):
    for by, value in seletores:
        try:
            for elemento in driver.find_elements(by, value):
                if elemento.is_displayed():
                    texto = elemento.text.strip()
                    if texto:
                        return texto
        except WebDriverException:
            return ""
    return ""


def selecionar_janela_com_url(driver, trecho_url):
    for handle in driver.window_handles:
        driver.switch_to.window(handle)
        if trecho_url.lower() in (driver.current_url or "").lower():
            return True
    return False


def normalizar_url(url):
    parsed = urlparse(url or "")
    hostname = (parsed.hostname or "").lower()
    path = (parsed.path or "").rstrip("/").lower()
    return hostname, path


def url_corresponde(url_atual, url_esperada):
    host_atual, path_atual = normalizar_url(url_atual)
    host_esperado, path_esperado = normalizar_url(url_esperada)
    return host_atual == host_esperado and path_atual == path_esperado


def criar_driver():
    options = webdriver.ChromeOptions()
    chrome_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
    if os.path.exists(chrome_path):
        options.binary_location = chrome_path
    options.add_argument("--headless")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1920,1080")
    driver = webdriver.Chrome(options=options)
    driver.set_page_load_timeout(30)
    return driver


def encerrar_driver(driver):
    if driver:
        try:
            driver.quit()
        except WebDriverException:
            pass


def validar_sessao_limpa(driver):
    try:
        cookies_limpos = len(driver.get_cookies()) == 0
        storage_limpo = driver.execute_script("return window.localStorage.length + window.sessionStorage.length;") == 0
        uma_janela = len(driver.window_handles) == 1
        return cookies_limpos and storage_limpo and uma_janela
    except WebDriverException:
        return False


def limpar_sessao_driver(driver):
    try:
        handles = driver.window_handles
        if not handles:
            return False
        principal = handles[0]
        for handle in handles[1:]:
            driver.switch_to.window(handle)
            driver.close()
        driver.switch_to.window(principal)
        driver.delete_all_cookies()
        driver.execute_script("window.localStorage.clear();")
        driver.execute_script("window.sessionStorage.clear();")
        driver.get("about:blank")
        driver.delete_all_cookies()
        return validar_sessao_limpa(driver)
    except WebDriverException:
        return False


def recriar_driver(driver):
    encerrar_driver(driver)
    return criar_driver()


def reiniciar_driver(driver, logger=None, linha=None, sistema=None, cpf=None, motivo="reinicializacao"):
    erro_encerramento = ""
    try:
        if driver:
            driver.quit()
    except Exception as exc:
        erro_encerramento = str(exc)
        if logger:
            log_event(
                logger,
                logging.WARNING,
                "AUTENTICACAO",
                linha,
                sistema,
                cpf,
                "reiniciar_driver",
                ERRO_SELENIUM,
                0,
                f"Falha ao encerrar driver antigo durante {motivo}: {erro_encerramento}",
            )
    novo_driver = None
    try:
        novo_driver = criar_driver()
        if logger:
            msg = "Driver reiniciado com sucesso."
            if erro_encerramento:
                msg += f" Encerramento anterior retornou erro: {erro_encerramento}"
            log_event(logger, logging.INFO, "AUTENTICACAO", linha, sistema, cpf, "reiniciar_driver", SUCESSO, 0, msg)
        return novo_driver
    except Exception as exc:
        if logger:
            log_event(
                logger,
                logging.ERROR,
                "AUTENTICACAO",
                linha,
                sistema,
                cpf,
                "reiniciar_driver",
                ERRO_SELENIUM,
                0,
                f"Nao foi possivel recriar o driver apos erro na linha {linha}: {exc}",
            )
        raise


def validar_login_ms(driver, cpf, senha):
    try:
        driver.get(LOGIN_URL_MS)
        cpf_input = aguardar_primeiro_elemento_valido(driver, SeletoresMS.CPF, timeout=10)
        senha_input = aguardar_primeiro_elemento_valido(driver, SeletoresMS.SENHA, timeout=10)
        botao = aguardar_primeiro_elemento_valido(driver, SeletoresMS.BOTAO_LOGIN, timeout=10, clicavel=True)
        cpf_input.clear()
        senha_input.clear()
        cpf_input.send_keys(cpf)
        senha_input.send_keys(senha)
        botao.click()

        try:
            WebDriverWait(driver, 10).until(
                lambda d: selecionar_janela_com_url(d, "home.aspx") or elemento_existe(d, SeletoresMS.MENSAGEM_ERRO)
            )
        except TimeoutException:
            return ResultadoAuth(TIMEOUT, resultado_planilha_por_classificacao(TIMEOUT), mensagem_para_planilha(TIMEOUT))

        mensagem = ""
        try:
            mensagem = limpar_mensagem_autenticacao(aguardar_primeiro_elemento_valido(driver, SeletoresMS.MENSAGEM_ERRO, timeout=2).text)
        except TimeoutException:
            mensagem = ""

        if mensagem:
            classificacao = classificar_mensagem_autenticacao(mensagem)
            if classificacao == INCONCLUSIVO:
                return ResultadoAuth(INCONCLUSIVO, RESULTADO_INCONCLUSIVO, mensagem_para_planilha(INCONCLUSIVO), mensagem)
            return ResultadoAuth(classificacao, resultado_planilha_por_classificacao(classificacao), mensagem_para_planilha(classificacao, mensagem), mensagem)

        url_home = "home.aspx" in driver.current_url.lower()
        formulario_ausente = not elemento_existe(driver, SeletoresMS.CPF) and not elemento_existe(driver, SeletoresMS.SENHA)
        if url_home and formulario_ausente:
            return ResultadoAuth(SUCESSO, RESULTADO_SIM, mensagem_para_planilha(SUCESSO))
        if url_home:
            return ResultadoAuth(INCONCLUSIVO, RESULTADO_INCONCLUSIVO, mensagem_para_planilha(INCONCLUSIVO), "URL de home sem segundo indicador comprovado.")
        return ResultadoAuth(PAGINA_INESPERADA, RESULTADO_ERRO, mensagem_para_planilha(PAGINA_INESPERADA), driver.current_url)
    except TimeoutException as exc:
        return ResultadoAuth(TIMEOUT, RESULTADO_ERRO, mensagem_para_planilha(TIMEOUT), str(exc))
    except NoSuchElementException as exc:
        return ResultadoAuth(ELEMENTO_NAO_ENCONTRADO, RESULTADO_ERRO, mensagem_para_planilha(ELEMENTO_NAO_ENCONTRADO), str(exc))
    except WebDriverException as exc:
        return ResultadoAuth(ERRO_SELENIUM, RESULTADO_ERRO, mensagem_para_planilha(ERRO_SELENIUM), str(exc))
    except Exception as exc:
        return ResultadoAuth(ERRO_INTERNO, RESULTADO_ERRO, mensagem_para_planilha(ERRO_INTERNO), str(exc))


def validar_login_ad(driver, cpf, senha):
    try:
        driver.get(LOGIN_URL_AD)
        dropdown = aguardar_primeiro_elemento_valido(driver, SeletoresAD.DROPDOWN_LOGIN, timeout=10, clicavel=True)
        dropdown.click()
        cpf_input = aguardar_primeiro_elemento_valido(driver, SeletoresAD.CPF, timeout=10)
        senha_input = aguardar_primeiro_elemento_valido(driver, SeletoresAD.SENHA, timeout=10)
        botao = aguardar_primeiro_elemento_valido(driver, SeletoresAD.BOTAO_AUTENTICAR, timeout=10, clicavel=True)
        cpf_input.clear()
        senha_input.clear()
        cpf_input.send_keys(cpf)
        senha_input.send_keys(senha)
        botao.click()

        try:
            WebDriverWait(driver, 10).until(
                lambda d: url_corresponde(d.current_url, SUCCESS_URL_AD) or bool(texto_visivel_primeiro_elemento(d, SeletoresAD.MENSAGEM_AUTENTICACAO))
            )
        except TimeoutException:
            return ResultadoAuth(TIMEOUT, RESULTADO_ERRO, mensagem_para_planilha(TIMEOUT), driver.current_url)

        mensagem = ""
        try:
            mensagem = limpar_mensagem_autenticacao(texto_visivel_primeiro_elemento(driver, SeletoresAD.MENSAGEM_AUTENTICACAO))
        except WebDriverException:
            mensagem = ""

        if mensagem:
            classificacao = classificar_mensagem_autenticacao(mensagem)
            if classificacao == INCONCLUSIVO:
                return ResultadoAuth(INCONCLUSIVO, RESULTADO_INCONCLUSIVO, mensagem_para_planilha(INCONCLUSIVO), mensagem)
            return ResultadoAuth(classificacao, resultado_planilha_por_classificacao(classificacao), mensagem_para_planilha(classificacao, mensagem), mensagem)

        url_sucesso = url_corresponde(driver.current_url, SUCCESS_URL_AD)
        formulario_ausente = not elemento_existe(driver, SeletoresAD.CPF) and not elemento_existe(driver, SeletoresAD.SENHA)
        if url_sucesso and formulario_ausente:
            return ResultadoAuth(SUCESSO, RESULTADO_SIM, mensagem_para_planilha(SUCESSO))
        if url_sucesso:
            return ResultadoAuth(INCONCLUSIVO, RESULTADO_INCONCLUSIVO, mensagem_para_planilha(INCONCLUSIVO), "URL de portal sem segundo indicador comprovado.")
        return ResultadoAuth(PAGINA_INESPERADA, RESULTADO_ERRO, mensagem_para_planilha(PAGINA_INESPERADA), f"URL final: {driver.current_url}")
    except TimeoutException as exc:
        return ResultadoAuth(TIMEOUT, RESULTADO_ERRO, mensagem_para_planilha(TIMEOUT), str(exc))
    except NoSuchElementException as exc:
        return ResultadoAuth(ELEMENTO_NAO_ENCONTRADO, RESULTADO_ERRO, mensagem_para_planilha(ELEMENTO_NAO_ENCONTRADO), str(exc))
    except WebDriverException as exc:
        return ResultadoAuth(ERRO_SELENIUM, RESULTADO_ERRO, mensagem_para_planilha(ERRO_SELENIUM), str(exc))
    except Exception as exc:
        return ResultadoAuth(ERRO_INTERNO, RESULTADO_ERRO, mensagem_para_planilha(ERRO_INTERNO), str(exc))


def atualizar_contador_auth(resumo, sistema, classificacao):
    chave = f"{sistema}_{classificacao}"
    resumo[chave] += 1


def deve_repetir_teste(resultado):
    return resultado.resultado_planilha in (RESULTADO_ERRO, RESULTADO_INCONCLUSIVO)


def mensagem_resultado_final_com_retentativa(primeira, segunda):
    detalhe = segunda.detalhe_tecnico or segunda.mensagem or primeira.detalhe_tecnico or primeira.mensagem
    if segunda.resultado_planilha == RESULTADO_SIM:
        return "Autenticacao realizada com sucesso apos reinicializacao do driver."
    if segunda.resultado_planilha == RESULTADO_ERRO:
        return f"Erro persistente apos reinicializacao do driver: {detalhe}"
    if segunda.resultado_planilha == RESULTADO_INCONCLUSIVO:
        return f"Resultado inconclusivo apos nova tentativa: {detalhe}"
    return segunda.mensagem


def executar_teste_autenticacao(driver, cpf, senha, funcao_teste):
    try:
        return funcao_teste(driver, cpf, senha)
    except Exception as exc:
        return ResultadoAuth(ERRO_INTERNO, RESULTADO_ERRO, mensagem_para_planilha(ERRO_INTERNO), str(exc))


def executar_com_uma_retentativa(sistema, linha, cpf, senha, driver, funcao_teste, logger, resumo):
    inicio = time.perf_counter()
    resultado_primeira = executar_teste_autenticacao(driver, cpf, senha, funcao_teste)
    mensagem_log = resultado_primeira.mensagem
    if resultado_primeira.detalhe_tecnico:
        mensagem_log = f"{mensagem_log} | detalhe={resultado_primeira.detalhe_tecnico}"
    log_event(
        logger,
        logging.INFO,
        "AUTENTICACAO",
        linha,
        sistema,
        cpf,
        f"validar_login_{sistema.lower()}",
        resultado_primeira.classificacao,
        time.perf_counter() - inicio,
        f"Primeira tentativa: {resultado_primeira.resultado_planilha} | {mensagem_log}",
    )

    if not deve_repetir_teste(resultado_primeira):
        return resultado_primeira, driver, False

    resumo["retentativas_auth"] += 1
    resumo[f"{sistema}_retentativas"] += 1
    mensagem_console = f"Erro inesperado na linha {linha}, CPF {cpf}, reiniciando driver e testando novamente"
    print(f"     {mensagem_console}")
    log_event(
        logger,
        logging.WARNING,
        "AUTENTICACAO",
        linha,
        sistema,
        cpf,
        "retentativa_auth",
        resultado_primeira.classificacao,
        0,
        (
            f"Erro inesperado na linha {linha}, CPF {mascarar_cpf(cpf)}, reiniciando driver e testando novamente | "
            f"primeira={resultado_primeira.resultado_planilha}; mensagem={mensagem_log}"
        ),
    )

    try:
        driver = reiniciar_driver(driver, logger=logger, linha=linha, sistema=sistema, cpf=cpf, motivo="retentativa")
        resumo["reinicializacoes_por_erro"] += 1
    except Exception as exc:
        resumo["erros_persistentes_auth"] += 1
        erro = ResultadoAuth(
            ERRO_SELENIUM,
            RESULTADO_ERRO,
            f"Nao foi possivel recriar o driver apos erro na linha {linha}: {exc}",
            str(exc),
        )
        log_event(logger, logging.ERROR, "AUTENTICACAO", linha, sistema, cpf, "retentativa_auth", erro.classificacao, 0, erro.mensagem)
        return erro, None, True

    inicio_segunda = time.perf_counter()
    resultado_segunda = executar_teste_autenticacao(driver, cpf, senha, funcao_teste)
    if resultado_segunda.resultado_planilha in (RESULTADO_SIM, RESULTADO_ERRO, RESULTADO_INCONCLUSIVO):
        resultado_segunda.mensagem = mensagem_resultado_final_com_retentativa(resultado_primeira, resultado_segunda)
    mensagem_log_segunda = resultado_segunda.mensagem
    if resultado_segunda.detalhe_tecnico:
        mensagem_log_segunda = f"{mensagem_log_segunda} | detalhe={resultado_segunda.detalhe_tecnico}"
    log_event(
        logger,
        logging.INFO,
        "AUTENTICACAO",
        linha,
        sistema,
        cpf,
        "retentativa_auth",
        resultado_segunda.classificacao,
        time.perf_counter() - inicio_segunda,
        (
            f"Segunda tentativa: {resultado_segunda.resultado_planilha} | {mensagem_log_segunda}; "
            f"resultado_final={resultado_segunda.resultado_planilha}"
        ),
    )
    if deve_repetir_teste(resultado_segunda):
        resumo["erros_persistentes_auth"] += 1
    else:
        resumo["erros_recuperados_auth"] += 1
    return resultado_segunda, driver, True


def chave_cache_autenticacao(cpf, senha, sistema):
    cpf_normalizado = normalizar_cpf_para_comparacao(cpf) or str(cpf or "").strip()
    return (sistema, cpf_normalizado, senha_para_chave(senha))


def copiar_resultado_auth(resultado):
    return ResultadoAuth(
        resultado.classificacao,
        resultado.resultado_planilha,
        resultado.mensagem,
        resultado.detalhe_tecnico,
    )


def resultado_auth_do_cache(resultado):
    resultado_cache = copiar_resultado_auth(resultado)
    if resultado_cache.resultado_planilha == RESULTADO_SIM:
        resultado_cache.mensagem = "cache: autenticação realizada com sucesso."
    elif resultado_cache.resultado_planilha == RESULTADO_NAO:
        resultado_cache.mensagem = "cache: " + resultado_cache.mensagem
    elif resultado_cache.resultado_planilha == RESULTADO_ERRO:
        resultado_cache.mensagem = "cache: " + resultado_cache.mensagem
    elif resultado_cache.resultado_planilha == RESULTADO_INCONCLUSIVO:
        resultado_cache.mensagem = "cache: " + resultado_cache.mensagem
    return resultado_cache


def obter_resultado_autenticacao_com_cache(cache_auth, sistema, linha, cpf, senha, driver, funcao_teste, logger, resumo):
    chave = chave_cache_autenticacao(cpf, senha, sistema)
    if chave in cache_auth:
        resumo["cache_auth_hits"] += 1
        resumo[f"{sistema}_cache_auth_hits"] += 1
        log_event(
            logger,
            logging.INFO,
            "AUTENTICACAO",
            linha,
            sistema,
            cpf,
            "cache_autenticacao",
            SUCESSO,
            0,
            "Resultado de autenticacao reutilizado do cache por CPF+senha nesta execucao.",
        )
        return resultado_auth_do_cache(cache_auth[chave]), driver, False, True
    resumo["cache_auth_misses"] += 1
    resultado, driver, reiniciou = executar_com_uma_retentativa(sistema, linha, cpf, senha, driver, funcao_teste, logger, resumo)
    cache_auth[chave] = copiar_resultado_auth(resultado)
    return resultado, driver, reiniciou, False


def registrar_resultado_autenticacao(ws, row, sistema, resultado, logger, resumo):
    if sistema == "MS":
        col_resultado = COL_RESULTADO_MS
        col_mensagem = COL_MSG_MS
    else:
        col_resultado = COL_RESULTADO_AD
        col_mensagem = COL_MSG_AD
    ws[f"{col_resultado}{row}"] = resultado.resultado_planilha
    ws[f"{col_mensagem}{row}"] = resultado.mensagem
    atualizar_contador_auth(resumo, sistema, resultado.classificacao)
    if resultado.classificacao != SUCESSO:
        adicionar_erro(resumo, "erros_autenticacao", row, f"{sistema} - {resultado.mensagem}")
        print(f"     ❌ {sistema}: {resultado.resultado_planilha} | {resultado.mensagem}")
    else:
        print(f"     ✅ {sistema}: SIM")


def resultado_erro_driver(mensagem, detalhe=""):
    return ResultadoAuth(ERRO_SELENIUM, RESULTADO_ERRO, mensagem, detalhe)


def tentar_criar_driver_para_continuar(logger, linha, sistema, cpf, resumo, motivo):
    try:
        driver = criar_driver()
        resumo["drivers_criados_auth"] += 1
        log_event(logger, logging.INFO, "AUTENTICACAO", linha, sistema, cpf, "criar_driver", SUCESSO, 0, f"Driver criado para continuar apos {motivo}.")
        return driver, None
    except Exception as exc:
        mensagem = f"Nao foi possivel criar driver para continuar apos {motivo} na linha {linha}: {exc}"
        log_event(logger, logging.ERROR, "AUTENTICACAO", linha, sistema, cpf, "criar_driver", ERRO_SELENIUM, 0, mensagem)
        return None, resultado_erro_driver(mensagem, str(exc))


def resultados_nt_em_ambos(ws, row):
    return ws[f"{COL_RESULTADO_MS}{row}"].value == RESULTADO_NT and ws[f"{COL_RESULTADO_AD}{row}"].value == RESULTADO_NT


def filtrar_linhas_nao_testaveis(ws, logger, resumo):
    removidas = []
    for row in range(ws.max_row, 1, -1):
        if resultados_nt_em_ambos(ws, row):
            removidas.append(row)
            ws.delete_rows(row, 1)
    removidas = sorted(removidas)
    resumo["linhas_removidas_nt"] = len(removidas)
    resumo["linhas_removidas_nt_numeros"] = removidas
    resumo["linhas_mantidas_pos_filtro"] = max(ws.max_row - 1, 0)
    mensagem = (
        f"Linhas removidas por ausencia de dados testaveis: {len(removidas)}; "
        f"Linhas originais removidas: {', '.join(map(str, removidas)) if removidas else '-'}; "
        "Regra aplicada: MS = N/T e AD = N/T"
    )
    log_event(logger, logging.INFO, "AUTENTICACAO", "-", "PLANILHA", None, "filtrar_linhas_nt", SUCESSO, 0, mensagem)
    return removidas


def processar_identidade(ws, cliente_api, logger, resumo):
    print(f"🔎 Iniciando resolução de identidade")
    print(f"📄 Entrada: {arquivo_excel}\n")
    for row in range(2, ws.max_row + 1):
        inicio = time.perf_counter()
        email_informado = str(ws[f"{COL_EMAIL_INFORMADO}{row}"].value or "").strip()
        documento_original = ws[f"{COL_DOCUMENTO_INFORMADO}{row}"].value
        ws[f"{COL_STATUS_PROCESSAMENTO}{row}"] = None
        ws[f"{COL_ORIGEM_CPF}{row}"] = None

        if not documento_original and not email_informado:
            ws[f"{COL_RESULTADO_MS}{row}"] = RESULTADO_NT
            ws[f"{COL_RESULTADO_AD}{row}"] = RESULTADO_NT
            ws[f"{COL_MSG_MS}{row}"] = "CPF/RE e e-mail ausentes."
            ws[f"{COL_MSG_AD}{row}"] = "CPF/RE e e-mail ausentes."
            ws[f"{COL_STATUS_PROCESSAMENTO}{row}"] = ERRO_DE_ENTRADA
            ws[f"{COL_ORIGEM_CPF}{row}"] = NAO_ENCONTRADO
            resumo["linhas_ignoradas"] += 1
            adicionar_erro(resumo, "erros_identidade", row, "Dados ausentes nas colunas A e C")
            print(f"  ❌ Linha {row}: dados ausentes")
            continue

        cpf = None
        origem = NAO_ENCONTRADO
        re_informado = None
        cpf_informado = None
        normalizado = normalizar_identificador(documento_original)
        if normalizado.status == SUCESSO and normalizado.valor.tipo == "CPF":
            cpf = normalizado.valor.valor
            cpf_informado = cpf
            origem = normalizado.valor.origem
            resumo["cpfs_validos_informados"] += 1
            if origem == CPF_RECONSTRUIDO_ZERO_INICIAL:
                resumo["cpfs_reconstruidos_zero_inicial"] += 1
        elif normalizado.status == SUCESSO and normalizado.valor.tipo == "RE":
            re_informado = normalizado.valor.valor
            resposta_cpf = cliente_api.buscar_cpf_por_re(normalizado.valor.valor, linha=row)
            if resposta_cpf.status == SUCESSO:
                cpf = resposta_cpf.valor
                origem = normalizado.valor.origem
                resumo["cpfs_encontrados_por_re"] += 1
            else:
                resumo["erros_api"] += 1 if resposta_cpf.status in CLASSIFICACOES_TECNICAS else 0
                ws[f"{COL_STATUS_PROCESSAMENTO}{row}"] = ERRO_TECNICO if resposta_cpf.status in CLASSIFICACOES_TECNICAS else None
                adicionar_erro(resumo, "erros_identidade", row, f"Erro ao buscar CPF pelo RE: {resposta_cpf.mensagem}")
                log_event(logger, logging.WARNING, "IDENTIDADE", row, "API", cpf, "buscar_cpf_por_re", resposta_cpf.status, time.perf_counter() - inicio, resposta_cpf.mensagem)
        elif normalizado.status not in (DADOS_AUSENTES,):
            ws[f"{COL_STATUS_PROCESSAMENTO}{row}"] = ERRO_DE_ENTRADA

        if cpf and email_informado:
            validacao_ad = buscar_usuario_ad_com_identificadores(email=email_informado, cpf=cpf, logger=logger, linha=row)
            if validacao_ad.status == AMBIGUO:
                ws[f"{COL_STATUS_PROCESSAMENTO}{row}"] = ERRO_DE_ENTRADA
                adicionar_erro(resumo, "erros_identidade", row, validacao_ad.mensagem)
                log_event(logger, logging.WARNING, "IDENTIDADE", row, "AD", cpf, "validar_conflito_ad", AMBIGUO, time.perf_counter() - inicio, validacao_ad.mensagem)
                cpf = None
                origem = AMBIGUO
            elif validacao_ad.status in CLASSIFICACOES_TECNICAS:
                log_event(logger, logging.WARNING, "IDENTIDADE", row, "AD", cpf, "validar_conflito_ad", validacao_ad.status, time.perf_counter() - inicio, validacao_ad.mensagem)

        if not cpf and email_informado:
            resposta_ad = buscar_usuario_ad_com_identificadores(email=email_informado, re_militar=re_informado, cpf=cpf_informado, logger=logger, linha=row)
            if resposta_ad.status == SUCESSO:
                usuario_ad = resposta_ad.valor
                if usuario_ad.cpf:
                    cpf = usuario_ad.cpf
                    origem = EMAIL_DSQUERY if not re_informado else RE_API
                    resumo["cpfs_encontrados_por_email"] += 1 if not re_informado else 0
                elif usuario_ad.re:
                    resposta_cpf = cliente_api.buscar_cpf_por_re(usuario_ad.re, linha=row)
                    if resposta_cpf.status == SUCESSO:
                        cpf = resposta_cpf.valor
                        origem = EMAIL_DSQUERY
                        resumo["cpfs_encontrados_por_email"] += 1
                    else:
                        resumo["erros_api"] += 1 if resposta_cpf.status in CLASSIFICACOES_TECNICAS else 0
                        adicionar_erro(resumo, "erros_identidade", row, f"Erro ao buscar CPF pelo RE do AD: {resposta_cpf.mensagem}")
                        log_event(logger, logging.WARNING, "IDENTIDADE", row, "API", cpf, "buscar_cpf_por_re_ad", resposta_cpf.status, time.perf_counter() - inicio, resposta_cpf.mensagem)
                else:
                    adicionar_erro(resumo, "erros_identidade", row, resposta_ad.mensagem)
            elif resposta_ad.status in (AMBIGUO, INCONCLUSIVO):
                adicionar_erro(resumo, "erros_identidade", row, resposta_ad.mensagem)
                log_event(logger, logging.WARNING, "IDENTIDADE", row, "AD", cpf, "resolver_identidade_ad", resposta_ad.status, time.perf_counter() - inicio, resposta_ad.mensagem)
            elif resposta_ad.status in CLASSIFICACOES_TECNICAS or resposta_ad.status == ACESSO_NEGADO:
                resumo["erros_dsquery"] += 1
                adicionar_erro(resumo, "erros_identidade", row, resposta_ad.mensagem)
                log_event(logger, logging.WARNING, "IDENTIDADE", row, "AD", cpf, "resolver_identidade_ad", resposta_ad.status, time.perf_counter() - inicio, resposta_ad.mensagem)

        if not cpf and re_informado and not email_informado:
            resposta_ad = buscar_usuario_ad_com_identificadores(re_militar=re_informado, logger=logger, linha=row)
            if resposta_ad.status == SUCESSO:
                usuario_ad = resposta_ad.valor
                if usuario_ad.cpf:
                    cpf = usuario_ad.cpf
                    origem = RE_API
                elif usuario_ad.re:
                    resposta_cpf = cliente_api.buscar_cpf_por_re(usuario_ad.re, linha=row)
                    if resposta_cpf.status == SUCESSO:
                        cpf = resposta_cpf.valor
                        origem = RE_API
            elif resposta_ad.status in CLASSIFICACOES_TECNICAS or resposta_ad.status == ACESSO_NEGADO:
                resumo["erros_dsquery"] += 1

        if not cpf:
            ws[f"{COL_CPF_RESOLVIDO}{row}"] = None
            ws[f"{COL_ORIGEM_CPF}{row}"] = NAO_ENCONTRADO
            ws[f"{COL_STATUS_PROCESSAMENTO}{row}"] = ws[f"{COL_STATUS_PROCESSAMENTO}{row}"].value or ERRO_DE_ENTRADA
            ws[f"{COL_RESULTADO_MS}{row}"] = RESULTADO_NT
            ws[f"{COL_RESULTADO_AD}{row}"] = RESULTADO_NT
            ws[f"{COL_MSG_MS}{row}"] = "CPF nao encontrado."
            ws[f"{COL_MSG_AD}{row}"] = "CPF nao encontrado."
            resumo["cpfs_nao_encontrados"] += 1
            adicionar_erro(resumo, "erros_identidade", row, "CPF não encontrado")
            print(f"  ❌ Linha {row}: CPF não encontrado")
            continue

        ws[f"{COL_CPF_RESOLVIDO}{row}"] = cpf
        ws[f"{COL_ORIGEM_CPF}{row}"] = origem
        nome = cliente_api.buscar_nome_por_cpf(cpf, linha=row)
        email = cliente_api.buscar_email_por_cpf(cpf, linha=row)
        if nome.status == SUCESSO:
            ws[f"{COL_NOME}{row}"] = nome.valor
        else:
            ws[f"{COL_NOME}{row}"] = "Nome nao encontrado pelo CPF"
            resumo["erros_api"] += 1 if nome.status in CLASSIFICACOES_TECNICAS else 0
            adicionar_erro(resumo, "erros_identidade", row, f"Erro ao consultar nome pelo CPF: {nome.mensagem}")
        if email.status == SUCESSO:
            ws[f"{COL_EMAIL_FUNCIONAL}{row}"] = email.valor
        else:
            ws[f"{COL_EMAIL_FUNCIONAL}{row}"] = "E-mail funcional nao encontrado pelo CPF"
            resumo["erros_api"] += 1 if email.status in CLASSIFICACOES_TECNICAS else 0
            adicionar_erro(resumo, "erros_identidade", row, f"Erro ao consultar e-mail funcional pelo CPF: {email.mensagem}")

        if nome.status == SUCESSO and email.status == SUCESSO:
            ws[f"{COL_STATUS_PROCESSAMENTO}{row}"] = CONCLUIDO
        else:
            ws[f"{COL_STATUS_PROCESSAMENTO}{row}"] = CONCLUIDO_COM_RESSALVA
        resumo["linhas_processadas"] += 1
        log_event(logger, logging.INFO, "IDENTIDADE", row, "API", cpf, "resolver_identidade", ws[f"{COL_STATUS_PROCESSAMENTO}{row}"].value, time.perf_counter() - inicio, "Identidade processada.")
        print(f"  ✅ Linha {row}: CPF {mascarar_cpf(cpf)} resolvido ({origem})")


def processar_autenticacoes(ws, logger, resumo):
    print("\nIniciando testes de autenticacao\n")
    driver = None
    linhas_com_driver = 0
    cache_auth = {}
    try:
        for row in range(2, ws.max_row + 1):
            driver_reiniciado_na_linha = False
            cpf = str(ws[f"{COL_CPF_RESOLVIDO}{row}"].value or "").strip()
            senha_original = ws[f"{COL_SENHA}{row}"].value
            senha = senha_para_chave(senha_original)
            ws[f"{COL_DATA_TESTE}{row}"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            cpf_ok = normalizar_cpf(cpf)
            if cpf_ok.status != SUCESSO or not senha:
                classificacao = DADOS_AUSENTES
                msg = "CPF ausente ou invalido." if cpf_ok.status != SUCESSO else "Senha ausente."
                ws[f"{COL_RESULTADO_MS}{row}"] = RESULTADO_NT
                ws[f"{COL_MSG_MS}{row}"] = msg
                ws[f"{COL_RESULTADO_AD}{row}"] = RESULTADO_NT
                ws[f"{COL_MSG_AD}{row}"] = msg
                atualizar_contador_auth(resumo, "MS", classificacao)
                atualizar_contador_auth(resumo, "AD", classificacao)
                adicionar_erro(resumo, "erros_autenticacao", row, f"MS/AD - {msg}")
                print(f"  Linha {row}: nao testada ({msg})")
                continue

            chave_ms = chave_cache_autenticacao(cpf, senha, "MS")
            chave_ad = chave_cache_autenticacao(cpf, senha, "AD")
            if chave_ms in cache_auth and chave_ad in cache_auth:
                resultado_ms, driver, _reiniciou_ms, _cache_ms = obter_resultado_autenticacao_com_cache(cache_auth, "MS", row, cpf, senha, driver, validar_login_ms, logger, resumo)
                registrar_resultado_autenticacao(ws, row, "MS", resultado_ms, logger, resumo)
                resultado_ad, driver, _reiniciou_ad, _cache_ad = obter_resultado_autenticacao_com_cache(cache_auth, "AD", row, cpf, senha, driver, validar_login_ad, logger, resumo)
                registrar_resultado_autenticacao(ws, row, "AD", resultado_ad, logger, resumo)
                print(f"  Linha {row}: resultados MS/AD reutilizados do cache | CPF {mascarar_cpf(cpf)}")
                continue

            if not driver:
                driver, erro_driver = tentar_criar_driver_para_continuar(logger, row, "SELENIUM", cpf, resumo, "inicio_da_linha")
                if erro_driver:
                    registrar_resultado_autenticacao(ws, row, "MS", erro_driver, logger, resumo)
                    registrar_resultado_autenticacao(ws, row, "AD", erro_driver, logger, resumo)
                    continue
                linhas_com_driver = 0
            elif linhas_com_driver >= 8:
                try:
                    driver = reiniciar_driver(driver, logger=logger, linha=row, sistema="SELENIUM", cpf=cpf, motivo="recriacao_periodica")
                except Exception as exc:
                    erro_driver = resultado_erro_driver(f"Nao foi possivel recriar driver periodicamente na linha {row}: {exc}", str(exc))
                    registrar_resultado_autenticacao(ws, row, "MS", erro_driver, logger, resumo)
                    registrar_resultado_autenticacao(ws, row, "AD", erro_driver, logger, resumo)
                    driver = None
                    continue
                linhas_com_driver = 0
                driver_reiniciado_na_linha = True
                resumo["drivers_criados_auth"] += 1
                resumo["reinicializacoes_periodicas"] += 1

            if not limpar_sessao_driver(driver):
                try:
                    driver = reiniciar_driver(driver, logger=logger, linha=row, sistema="SELENIUM", cpf=cpf, motivo="sessao_suja")
                except Exception as exc:
                    erro_driver = resultado_erro_driver(f"Nao foi possivel recriar driver apos sessao suja na linha {row}: {exc}", str(exc))
                    registrar_resultado_autenticacao(ws, row, "MS", erro_driver, logger, resumo)
                    registrar_resultado_autenticacao(ws, row, "AD", erro_driver, logger, resumo)
                    driver = None
                    continue
                linhas_com_driver = 0
                driver_reiniciado_na_linha = True
                resumo["drivers_criados_auth"] += 1
                resumo["reinicializacoes_por_erro"] += 1

            print(f"  Linha {row}: testando MS | CPF {mascarar_cpf(cpf)}")
            resultado_ms, driver, reiniciou_ms, cache_ms = obter_resultado_autenticacao_com_cache(cache_auth, "MS", row, cpf, senha, driver, validar_login_ms, logger, resumo)
            if reiniciou_ms:
                linhas_com_driver = 0
                driver_reiniciado_na_linha = True
                if driver:
                    resumo["drivers_criados_auth"] += 1
            registrar_resultado_autenticacao(ws, row, "MS", resultado_ms, logger, resumo)
            if cache_ms:
                print(f"     MS reutilizado do cache")
            if not driver:
                driver, erro_driver = tentar_criar_driver_para_continuar(logger, row, "AD", cpf, resumo, "falha_na_retentativa_ms")
                if erro_driver:
                    registrar_resultado_autenticacao(ws, row, "AD", erro_driver, logger, resumo)
                    continue

            if not limpar_sessao_driver(driver):
                try:
                    driver = reiniciar_driver(driver, logger=logger, linha=row, sistema="SELENIUM", cpf=cpf, motivo="limpeza_apos_ms")
                except Exception as exc:
                    erro_driver = resultado_erro_driver(f"Nao foi possivel recriar driver apos MS na linha {row}: {exc}", str(exc))
                    registrar_resultado_autenticacao(ws, row, "AD", erro_driver, logger, resumo)
                    driver = None
                    continue
                linhas_com_driver = 0
                driver_reiniciado_na_linha = True
                resumo["drivers_criados_auth"] += 1
                resumo["reinicializacoes_por_erro"] += 1

            print(f"  Linha {row}: testando AD | CPF {mascarar_cpf(cpf)}")
            resultado_ad, driver, reiniciou_ad, cache_ad = obter_resultado_autenticacao_com_cache(cache_auth, "AD", row, cpf, senha, driver, validar_login_ad, logger, resumo)
            if reiniciou_ad:
                linhas_com_driver = 0
                driver_reiniciado_na_linha = True
                if driver:
                    resumo["drivers_criados_auth"] += 1
            registrar_resultado_autenticacao(ws, row, "AD", resultado_ad, logger, resumo)
            if cache_ad:
                print(f"     AD reutilizado do cache")
            if not driver:
                continue

            if not limpar_sessao_driver(driver):
                try:
                    driver = reiniciar_driver(driver, logger=logger, linha=row, sistema="SELENIUM", cpf=cpf, motivo="limpeza_apos_ad")
                except Exception as exc:
                    log_event(logger, logging.ERROR, "AUTENTICACAO", row, "SELENIUM", cpf, "limpeza_apos_ad", ERRO_SELENIUM, 0, f"Nao foi possivel recriar driver apos AD na linha {row}: {exc}")
                    driver = None
                    continue
                linhas_com_driver = 0
                driver_reiniciado_na_linha = True
                resumo["drivers_criados_auth"] += 1
                resumo["reinicializacoes_por_erro"] += 1

            linhas_com_driver += 1
            if driver_reiniciado_na_linha:
                log_event(
                    logger,
                    logging.INFO,
                    "AUTENTICACAO",
                    row,
                    "SELENIUM",
                    cpf,
                    "contador_driver",
                    SUCESSO,
                    0,
                    f"Linhas desde a ultima reinicializacao: {linhas_com_driver}",
                )
    finally:
        encerrar_driver(driver)
    filtrar_linhas_nao_testaveis(ws, logger, resumo)

def remover_senhas(ws):
    for row in range(2, ws.max_row + 1):
        if ws[f"{COL_SENHA}{row}"].value:
            ws[f"{COL_SENHA}{row}"] = "REMOVIDA_APOS_PROCESSAMENTO"


def criar_aba_resumo(wb, resumo):
    if "Resumo" in wb.sheetnames:
        del wb["Resumo"]
    ws = wb.create_sheet("Resumo")
    ws["A1"] = "Campo"
    ws["B1"] = "Valor"
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    linhas = [
        ("ID da execucao", resumo["execution_id"]),
        ("Arquivo de entrada", resumo["arquivo_entrada"]),
        ("Arquivo final", resumo["arquivo_final"]),
        ("Inicio", resumo["inicio"].strftime("%Y-%m-%d %H:%M:%S")),
        ("Termino", resumo["termino"].strftime("%Y-%m-%d %H:%M:%S")),
        ("Duracao total (s)", round((resumo["termino"] - resumo["inicio"]).total_seconds(), 2)),
        ("Total de linhas", resumo["total_linhas"]),
        ("Linhas processadas", resumo["linhas_processadas"]),
        ("Linhas ignoradas", resumo["linhas_ignoradas"]),
        ("CPFs validos informados", resumo["cpfs_validos_informados"]),
        ("CPFs reconstruidos com zero inicial", resumo["cpfs_reconstruidos_zero_inicial"]),
        ("CPFs encontrados por RE", resumo["cpfs_encontrados_por_re"]),
        ("CPFs encontrados por e-mail", resumo["cpfs_encontrados_por_email"]),
        ("CPFs nao encontrados", resumo["cpfs_nao_encontrados"]),
        ("Erros de API", resumo["erros_api"]),
        ("Erros de dsquery", resumo["erros_dsquery"]),
    ]
    for sistema in ("MS", "AD"):
        for classificacao in [SUCESSO, CREDENCIAL_INVALIDA, CONTA_BLOQUEADA, CONTA_DESATIVADA, SENHA_EXPIRADA, NAO_TESTADO, INCONCLUSIVO, TIMEOUT, SISTEMA_INDISPONIVEL, ERRO_SELENIUM, ERRO_INTERNO, DADOS_AUSENTES]:
            linhas.append((f"{sistema} - {classificacao}", resumo.get(f"{sistema}_{classificacao}", 0)))
    for idx, (campo, valor) in enumerate(linhas, start=2):
        ws[f"A{idx}"] = campo
        ws[f"B{idx}"] = valor
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 80


def salvar_atomicamente(wb, destino):
    diretorio = os.path.dirname(destino)
    temp_path = None
    try:
        fd, temp_path = tempfile.mkstemp(prefix=".credenciais_", suffix=".xlsx", dir=diretorio)
        os.close(fd)
        wb.save(temp_path)
        if not os.path.exists(temp_path) or os.path.getsize(temp_path) == 0:
            raise RuntimeError("Arquivo temporario nao foi criado corretamente.")
        teste = load_workbook(temp_path, read_only=True)
        teste.close()
        os.replace(temp_path, destino)
    except Exception:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass
        raise


def obter_competencia_atual(data_execucao):
    mes = MESES_PT[data_execucao.month]
    ano = f"{data_execucao.year % 100:02d}"
    return f"{mes}{ano}"


def obter_dia_semana_portugues(data_execucao):
    return DIAS_SEMANA_PT[data_execucao.weekday()]


def normalizar_cpf_para_comparacao(valor):
    resultado = normalizar_cpf(valor)
    if resultado.status == SUCESSO:
        return resultado.valor
    return None


def resultado_eh_sim(valor):
    return str(valor or "").strip().upper() == RESULTADO_SIM


def extrair_registros_mensais(ws_processado, sistema):
    col_resultado = COL_RESULTADO_AD if sistema == "AD" else COL_RESULTADO_MS
    registros = []
    for row in range(2, ws_processado.max_row + 1):
        if not resultado_eh_sim(ws_processado[f"{col_resultado}{row}"].value):
            continue
        registros.append({
            "nome": ws_processado[f"{COL_NOME}{row}"].value,
            "email": ws_processado[f"{COL_EMAIL_FUNCIONAL}{row}"].value,
            "cpf": ws_processado[f"{COL_CPF_RESOLVIDO}{row}"].value,
            "senha": ws_processado[f"{COL_SENHA}{row}"].value,
            "ms": ws_processado[f"{COL_RESULTADO_MS}{row}"].value,
            "ad": ws_processado[f"{COL_RESULTADO_AD}{row}"].value,
            "msg_ms": ws_processado[f"{COL_MSG_MS}{row}"].value,
            "msg_ad": ws_processado[f"{COL_MSG_AD}{row}"].value,
            "url_vazamento": ws_processado[f"{COL_URL_VAZAMENTO}{row}"].value,
        })
    return registros


def criar_workbook_mensal():
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultado"
    for col, header in enumerate(HEADERS_MENSAL, start=1):
        cell = ws.cell(1, col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    aplicar_formatacao_mensal(ws)
    return wb


def validar_layout_mensal(ws):
    headers_existentes = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    while headers_existentes and headers_existentes[-1] is None:
        headers_existentes.pop()
    headers_migrados = list(headers_existentes)
    if len(headers_migrados) >= 9:
        header_coluna_contagem = str(headers_migrados[8] or "").lower()
        if headers_migrados[8] != HEADER_CONTAGEM_IDENTIFICACOES and "identificado" in header_coluna_contagem:
            headers_migrados[8] = HEADER_CONTAGEM_IDENTIFICACOES
    if headers_migrados == HEADERS_MENSAL:
        for col, header in enumerate(HEADERS_MENSAL, start=1):
            ws.cell(1, col).value = header
        return
    if headers_migrados == HEADERS_MENSAL[:len(headers_migrados)]:
        for col in range(1, len(HEADERS_MENSAL) + 1):
            cell = ws.cell(1, col)
            cell.value = HEADERS_MENSAL[col - 1]
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        return
    raise ValueError("Layout mensal existente incompat?vel; arquivo preservado sem altera??o.")

def carregar_ou_criar_arquivo_mensal(caminho):
    if os.path.exists(caminho):
        wb = load_workbook(caminho)
        ws = wb.active
        validar_layout_mensal(ws)
        return wb, ws, False
    wb = criar_workbook_mensal()
    return wb, wb.active, True


def obter_contagem_identificacoes(valor):
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try:
            return max(int(valor), 1)
        except (TypeError, ValueError):
            return 1
    texto = str(valor or "").strip()
    if texto.isdigit():
        return max(int(texto), 1)
    return 1


def indexar_cpfs_existentes(ws):
    indice = {}
    for row in range(2, ws.max_row + 1):
        cpf = normalizar_cpf_para_comparacao(ws.cell(row, 3).value)
        if not cpf:
            continue
        primeira = ws.cell(row, 10).value
        if not primeira:
            primeira = ws.cell(row, 11).value
        contagem = obter_contagem_identificacoes(ws.cell(row, 9).value)
        if cpf not in indice:
            indice[cpf] = {
                "row": row,
                "primeira": primeira,
                "contagem": contagem,
            }
        else:
            indice[cpf]["contagem"] += contagem
            if not indice[cpf]["primeira"] and primeira:
                indice[cpf]["primeira"] = primeira
    return indice


def aplicar_formatacao_mensal(ws):
    success_fill = PatternFill("solid", fgColor="C6EFCE")
    failure_fill = PatternFill("solid", fgColor="FFC7CE")
    error_fill = PatternFill("solid", fgColor="F4B183")
    inconclusive_fill = PatternFill("solid", fgColor="FFEB9C")
    nt_fill = PatternFill("solid", fgColor="D9D9D9")
    resultado_fills = {
        RESULTADO_SIM: success_fill,
        RESULTADO_NAO: failure_fill,
        RESULTADO_ERRO: error_fill,
        RESULTADO_INCONCLUSIVO: inconclusive_fill,
        RESULTADO_NT: nt_fill,
    }
    for col in range(1, len(HEADERS_MENSAL) + 1):
        cell = ws.cell(1, col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(2, ws.max_row + 1):
        for col in (5, 6):
            cell = ws.cell(row, col)
            if cell.value in resultado_fills:
                cell.fill = resultado_fills[cell.value]
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in (3, 9, 10, 11, 12):
            ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center")
        for col in (7, 8, 13):
            ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (10, 11):
            ws.cell(row, col).number_format = "dd/mm/yyyy hh:mm:ss"
    larguras = {
        "A": 34,
        "B": 38,
        "C": 16,
        "D": 28,
        "E": 12,
        "F": 12,
        "G": 58,
        "H": 58,
        "I": 24,
        "J": 24,
        "K": 24,
        "L": 18,
        "M": 58,
    }
    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura
    ws.auto_filter.ref = f"A1:M{max(ws.max_row, 1)}"
    ws.freeze_panes = "A2"


def adicionar_registros_mensais(ws, registros, data_execucao):
    indice_cpfs = indexar_cpfs_existentes(ws)
    adicionados = 0
    duplicados = 0
    dia_semana = obter_dia_semana_portugues(data_execucao)
    for registro in registros:
        cpf_normalizado = normalizar_cpf_para_comparacao(registro["cpf"])
        primeira_identificacao = data_execucao
        contagem_identificacoes = 1
        linha_nova = True
        if cpf_normalizado and cpf_normalizado in indice_cpfs:
            duplicados += 1
            info_cpf = indice_cpfs[cpf_normalizado]
            row = info_cpf["row"]
            primeira_identificacao = info_cpf["primeira"] or data_execucao
            contagem_identificacoes = info_cpf["contagem"] + 1
            linha_nova = False
        else:
            row = ws.max_row + 1
        valores = [
            registro["nome"],
            registro["email"],
            registro["cpf"],
            registro["senha"],
            registro["ms"],
            registro["ad"],
            registro["msg_ms"],
            registro["msg_ad"],
            contagem_identificacoes,
            primeira_identificacao,
            data_execucao,
            dia_semana,
            registro["url_vazamento"],
        ]
        for col, valor in enumerate(valores, start=1):
            ws.cell(row, col).value = valor
        if cpf_normalizado:
            indice_cpfs[cpf_normalizado] = {
                "row": row,
                "primeira": primeira_identificacao,
                "contagem": contagem_identificacoes,
            }
        if linha_nova:
            adicionados += 1
    aplicar_formatacao_mensal(ws)
    return adicionados, duplicados


def salvar_workbook_mensal_atomicamente(wb, destino, linhas_antes, novas_linhas):
    diretorio = os.path.dirname(destino)
    temp_path = None
    try:
        fd, temp_path = tempfile.mkstemp(prefix=".mensal_", suffix=".xlsx", dir=diretorio)
        os.close(fd)
        wb.save(temp_path)
        teste = load_workbook(temp_path, read_only=True)
        ws_teste = teste.active
        headers = [ws_teste.cell(1, col).value for col in range(1, len(HEADERS_MENSAL) + 1)]
        if headers != HEADERS_MENSAL:
            teste.close()
            raise RuntimeError("Cabeçalhos do arquivo mensal ficaram inválidos.")
        if ws_teste.max_column != len(HEADERS_MENSAL):
            teste.close()
            raise RuntimeError("Quantidade de colunas do arquivo mensal ficou inválida.")
        if ws_teste.max_row != linhas_antes + novas_linhas:
            teste.close()
            raise RuntimeError("Quantidade de linhas do arquivo mensal não confere.")
        teste.close()
        os.replace(temp_path, destino)
    except PermissionError:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass
        raise PermissionError(f"Não foi possível atualizar {os.path.basename(destino)} porque o arquivo está aberto ou sem permissão de escrita.")
    except Exception:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass
        raise


def atualizar_arquivo_mensal(sistema, registros, data_execucao):
    competencia = obter_competencia_atual(data_execucao)
    nome_arquivo = f"Credenciais_{sistema}_{competencia}.xlsx"
    caminho = os.path.join(script_dir, nome_arquivo)
    resultado = {
        "sistema": sistema,
        "arquivo": caminho,
        "encontrados": len(registros),
        "adicionados": 0,
        "duplicados": 0,
        "erro": "",
    }
    if not registros:
        resultado["arquivo"] = caminho if os.path.exists(caminho) else ""
        return resultado
    try:
        wb, ws, _criado = carregar_ou_criar_arquivo_mensal(caminho)
        linhas_antes = ws.max_row
        adicionados, duplicados = adicionar_registros_mensais(ws, registros, data_execucao)
        salvar_workbook_mensal_atomicamente(wb, caminho, linhas_antes, adicionados)
        resultado["adicionados"] = adicionados
        resultado["duplicados"] = duplicados
    except Exception as exc:
        resultado["erro"] = str(exc)
    return resultado


def escrever_log_resumido(caminho, titulo_erros, erros, linhas_sucesso, linhas_falha, resultado_final, consolidacao_mensal=None):
    with open(caminho, "w", encoding="utf-8") as log:
        log.write(f"=== NOVA EXECUÇÃO: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} ===\n")
        log.write(f"{titulo_erros}\n")
        if erros:
            for linha, erro in erros:
                log.write(f"Linha {linha} - {erro}\n")
        else:
            log.write("Nenhum erro encontrado.\n")
        log.write(f"\n✔️ Registros com sucesso: {linhas_sucesso}\n")
        log.write(f"❌ Registros com falha: {linhas_falha}\n")
        log.write(f"📁 Planilha gerada: {resultado_final}\n")
        if isinstance(consolidacao_mensal, dict) and consolidacao_mensal.get("auditoria_nt"):
            auditoria = consolidacao_mensal["auditoria_nt"]
            removidas = auditoria.get("linhas_removidas", [])
            log.write("\nAuditoria de remocao N/T:\n")
            log.write(f"Total de linhas processadas: {auditoria.get('total_linhas', 0)}\n")
            log.write(f"Linhas removidas por ausencia de dados testaveis: {len(removidas)}\n")
            log.write(f"Linhas mantidas: {auditoria.get('linhas_mantidas', 0)}\n")
            log.write(f"Linhas originais removidas: {', '.join(map(str, removidas)) if removidas else '-'}\n")
            log.write("Regra aplicada: MS = N/T e AD = N/T\n")
        if isinstance(consolidacao_mensal, dict) and consolidacao_mensal.get("arquivo"):
            log.write("\nResumo mensal consolidado:\n")
            log.write(f"Arquivo mensal: {consolidacao_mensal.get('arquivo')}\n")
            log.write(f"Situacao: {consolidacao_mensal.get('situacao', '')}\n")
            log.write(f"Abas criadas: {', '.join(consolidacao_mensal.get('abas_criadas', [])) or '-'}\n")
            log.write(f"Registros recebidos: {consolidacao_mensal.get('registros_recebidos', 0)}\n")
            log.write(f"Registros positivos no AD: {consolidacao_mensal.get('positivos_ad', 0)}\n")
            log.write(f"Registros positivos no MS: {consolidacao_mensal.get('positivos_ms', 0)}\n")
            log.write(f"Registros enviados para Total: {consolidacao_mensal.get('total_enviados', 0)}\n")
            for sistema in ("AD", "MS", "Total"):
                info = consolidacao_mensal.get(sistema, {})
                if not info:
                    continue
                log.write(
                    f"{sistema}: novos={info.get('novos', 0)}, atualizados={info.get('atualizados', 0)}, "
                    f"contadores_incrementados={info.get('contadores_incrementados', 0)}, "
                    f"datas_acrescentadas={info.get('datas_acrescentadas', 0)}, "
                    f"resultados_atualizados={info.get('resultados_atualizados', 0)}, "
                    f"urls_acrescentadas={info.get('urls_acrescentadas', 0)}\n"
                )
            migracao = consolidacao_mensal.get("migracao", [])
            log.write(f"Migracao de arquivos antigos: {', '.join(migracao) if migracao else '-'}\n")
            if consolidacao_mensal.get("erro"):
                log.write(f"Erro na consolidacao mensal: {consolidacao_mensal['erro']}\n")
        if consolidacao_mensal:
            log.write(f"\nArquivo individual gerado:\n{resultado_final}\n")
            for sistema in ("AD", "MS"):
                info = consolidacao_mensal.get(sistema, {})
                if info.get("erro"):
                    log.write(f"\nFalha ao atualizar arquivo mensal do {sistema}: {info['erro']}\n")
                    continue
                if info.get("encontrados", 0) == 0:
                    log.write(f"\nNenhum resultado SIM encontrado para o {sistema} nesta execução.\n")
                    continue
                log.write(f"\nArquivo mensal {sistema}:\n{info.get('arquivo', '')}\n")
                log.write(f"Novas linhas adicionadas ao arquivo {sistema}:\n{info.get('adicionados', 0)}\n")
                log.write(f"Registros atualizados por identificacao repetida no {sistema}:\n{info.get('duplicados', 0)}\n")


def formatar_data_hora_local(valor):
    if not isinstance(valor, datetime):
        valor = datetime.now()
    return valor.strftime("%d/%m/%Y %H:%M:%S")


def formatar_duracao(inicio, termino):
    try:
        segundos = int(max((termino - inicio).total_seconds(), 0))
        horas, resto = divmod(segundos, 3600)
        minutos, segundos = divmod(resto, 60)
        return f"{horas:02d}:{minutos:02d}:{segundos:02d}"
    except Exception:
        return "n\u00e3o calculada"


def anexar_bloco_log(caminho, bloco):
    try:
        with open(caminho, "a", encoding="utf-8", newline="\n") as arquivo:
            arquivo.write(bloco)
            if not bloco.endswith("\n"):
                arquivo.write("\n")
    except Exception as exc:
        mensagem = f"N\u00e3o foi poss\u00edvel atualizar {os.path.basename(caminho)}: {exc}"
        print(mensagem)
        if LOG_CONTEXT:
            LOG_CONTEXT.falhas_log.append(mensagem)


def status_final_execucao(resumo):
    if resumo.get("erro_critico"):
        return "\u274c Execu\u00e7\u00e3o interrompida por erro cr\u00edtico"
    if LOG_CONTEXT and LOG_CONTEXT.eventos:
        return "\u26a0\ufe0f Execu\u00e7\u00e3o conclu\u00edda com ocorr\u00eancias"
    if resumo.get("erros_identidade") or resumo.get("erros_autenticacao"):
        return "\u26a0\ufe0f Execu\u00e7\u00e3o conclu\u00edda com ocorr\u00eancias"
    return "\u2705 Execu\u00e7\u00e3o conclu\u00edda com sucesso"


def nome_arquivo_ou_falha(caminho):
    if not caminho:
        return "falha na gera\u00e7\u00e3o"
    return os.path.basename(str(caminho))


def montar_bloco_log_execucao(resumo, resultado_final, consolidacao_mensal=None):
    contexto = LOG_CONTEXT or iniciar_contexto_logs(resumo.get("execution_id", datetime.now().strftime("%Y%m%d-%H%M%S")), resumo.get("inicio", datetime.now()))
    termino = resumo.get("termino") if isinstance(resumo.get("termino"), datetime) else datetime.now()
    mensal_erro = consolidacao_mensal.get("erro") if isinstance(consolidacao_mensal, dict) else ""
    mensal_nome = nome_arquivo_ou_falha(consolidacao_mensal.get("arquivo") if isinstance(consolidacao_mensal, dict) else "")
    atualizacao_mensal = "n\u00e3o conclu\u00edda \u2014 consulte Logs_auditoria.txt" if mensal_erro else "conclu\u00edda com sucesso"
    linhas = [
        "============================================================",
        f"\U0001f680 NOVA EXECU\u00c7\u00c3O: {formatar_data_hora_local(contexto.inicio)}",
        "============================================================",
        "",
        f"Identificador da execu\u00e7\u00e3o: {contexto.execution_id}",
        "",
        "\U0001f5a5\ufe0f IDENTIFICA\u00c7\u00c3O",
        f"Nome da m\u00e1quina: {contexto.maquina}",
        f"IP da m\u00e1quina: {contexto.ip}",
        f"Usu\u00e1rio logado: {contexto.usuario}",
        "",
        "\U0001f4ca RESUMO DO PROCESSAMENTO",
        f"Total de linhas recebidas: {resumo.get('total_linhas', 0)}",
        f"Total de linhas processadas: {resumo.get('linhas_processadas', 0)}",
        f"Linhas removidas por aus\u00eancia de dados test\u00e1veis: {resumo.get('linhas_removidas_nt', 0)}",
        f"Credenciais com acesso ao AD: {resumo.get('AD_SUCESSO', 0)}",
        f"Credenciais com acesso ao MS: {resumo.get('MS_SUCESSO', 0)}",
        "",
        "\U0001f4c4 ARQUIVOS",
        f"Planilha de resultado gerada: {nome_arquivo_ou_falha(resultado_final)}",
        f"Planilha mensal: {mensal_nome}",
        f"Atualiza\u00e7\u00e3o mensal: {atualizacao_mensal}",
        "",
        status_final_execucao(resumo),
        f"\u2705 Execu\u00e7\u00e3o finalizada em: {formatar_data_hora_local(termino)}",
        f"\u23f1\ufe0f Dura\u00e7\u00e3o total: {formatar_duracao(contexto.inicio, termino)}",
        "",
        "Observa\u00e7\u00e3o: os arquivos antigos Logs_BuscarNome.txt e Logs_testarCredenciais.txt n\u00e3o s\u00e3o mais utilizados.",
        "",
    ]
    return "\n".join(linhas)


def eventos_por_categoria():
    categorias = {"identidade": [], "credenciais": [], "planilha": [], "outros": []}
    if LOG_CONTEXT:
        for evento in LOG_CONTEXT.eventos:
            categorias.setdefault(evento.categoria, []).append(evento)
    return categorias


def titulo_categoria_auditoria(categoria):
    return {
        "identidade": "\U0001f50e ERROS DE CONSULTA DE IDENTIDADE",
        "credenciais": "\U0001f510 ERROS DE TESTE DE CREDENCIAIS",
        "planilha": "\U0001f4ca ERROS DE PLANILHA E CONSOLIDA\u00c7\u00c3O",
        "outros": "\u2699\ufe0f OUTROS ERROS T\u00c9CNICOS",
    }[categoria]


def renderizar_evento_auditoria(evento, indice):
    linhas = [
        f"[Erro {indice}]",
        f"Data/hora: {formatar_data_hora_local(evento.timestamp)}",
        f"Linha original: {evento.linha}",
    ]
    if evento.sistema and evento.sistema != "-":
        linhas.append(f"Sistema: {evento.sistema}")
    if evento.cpf and evento.cpf != "-":
        linhas.append(f"CPF: {evento.cpf}")
    if evento.re_militar:
        linhas.append(f"RE: {evento.re_militar}")
    if evento.email:
        linhas.append(f"E-mail: {evento.email}")
    linhas.extend([
        f"Etapa: {evento.operacao}",
        f"Tipo do erro: {evento.status}",
        f"Mensagem: {evento.mensagem}",
        "",
    ])
    return linhas


def montar_bloco_log_auditoria(resumo):
    contexto = LOG_CONTEXT or iniciar_contexto_logs(resumo.get("execution_id", datetime.now().strftime("%Y%m%d-%H%M%S")), resumo.get("inicio", datetime.now()))
    categorias = eventos_por_categoria()
    total = sum(len(v) for v in categorias.values())
    linhas = [
        "============================================================",
        f"\U0001f6e1\ufe0f AUDITORIA DA EXECU\u00c7\u00c3O: {formatar_data_hora_local(contexto.inicio)}",
        "============================================================",
        "",
        f"Identificador da execu\u00e7\u00e3o: {contexto.execution_id}",
        f"M\u00e1quina: {contexto.maquina}",
        f"Usu\u00e1rio: {contexto.usuario}",
        "",
    ]
    if total == 0:
        linhas.extend([
            "\u2705 Nenhum erro t\u00e9cnico ou de processamento foi registrado.",
            "============================================================",
            "",
        ])
        return "\n".join(linhas)
    for categoria in ("identidade", "credenciais", "planilha", "outros"):
        eventos = categorias.get(categoria, [])
        linhas.extend([
            "------------------------------------------------------------",
            titulo_categoria_auditoria(categoria),
            "------------------------------------------------------------",
            "",
        ])
        if not eventos:
            linhas.extend(["Nenhum erro registrado.", ""])
            continue
        for idx, evento in enumerate(eventos, start=1):
            linhas.extend(renderizar_evento_auditoria(evento, idx))
    linhas.extend([
        "------------------------------------------------------------",
        "\U0001f4cc RESUMO DA AUDITORIA",
        "------------------------------------------------------------",
        "",
        f"Erros de consulta de identidade: {len(categorias.get('identidade', []))}",
        f"Erros de teste de credenciais: {len(categorias.get('credenciais', []))}",
        f"Erros de planilha e consolida\u00e7\u00e3o: {len(categorias.get('planilha', []))}",
        f"Outros erros t\u00e9cnicos: {len(categorias.get('outros', []))}",
        f"Total de ocorr\u00eancias registradas: {total}",
        "",
    ])
    return "\n".join(linhas)


def escrever_logs_finais(resumo, resultado_final, consolidacao_mensal=None):
    if consolidacao_mensal and consolidacao_mensal.get("erro"):
        registrar_evento_auditoria(
            "planilha",
            linha="-",
            sistema="MENSAL",
            operacao="atualizacao_mensal",
            status=ERRO_TECNICO,
            mensagem=consolidacao_mensal["erro"],
        )
    if resumo.get("erro_critico"):
        registrar_evento_auditoria(
            "outros",
            linha="-",
            sistema="APLICACAO",
            operacao="main",
            status=ERRO_INTERNO,
            mensagem=resumo["erro_critico"],
        )
    anexar_bloco_log(os.path.join(script_dir, "Logs_execucao.txt"), montar_bloco_log_execucao(resumo, resultado_final, consolidacao_mensal))
    anexar_bloco_log(os.path.join(script_dir, "Logs_auditoria.txt"), montar_bloco_log_auditoria(resumo))

def resultado_consolidavel_total(valor):
    return str(valor or "").strip().upper() in (RESULTADO_SIM, RESULTADO_NAO, RESULTADO_ERRO, RESULTADO_INCONCLUSIVO)


def valor_util_mensal(valor):
    if valor is None:
        return ""
    texto = str(valor).strip()
    if not texto:
        return ""
    texto_normalizado = texto_sem_acentos(texto).lower()
    genericos = (
        "nome nao encontrado",
        "e-mail funcional nao encontrado",
        "email funcional nao encontrado",
        "cpf nao encontrado",
        "nao encontrado",
        "none",
        "nan",
    )
    if any(item in texto_normalizado for item in genericos):
        return ""
    return texto


def normalizar_url_vazamento(valor):
    texto = valor_util_mensal(valor)
    if not texto:
        return ""
    parsed = urlparse(texto)
    if parsed.scheme and parsed.netloc:
        return parsed._replace(query="", fragment="").geturl()
    return re.split(r"[?#]", texto, maxsplit=1)[0].strip()


def formatar_data_identificacao(data_execucao):
    return f"{data_execucao.day:02d}{MESES_PT[data_execucao.month]}"


def resultado_mensal_acesso(valor):
    texto = str(valor or "").strip().upper()
    if texto == RESULTADO_SIM:
        return RESULTADO_SIM
    if texto == RESULTADO_NAO:
        return RESULTADO_NAO
    return ""


def registro_mensal_processavel(ws_processado, row):
    acesso_ad = resultado_mensal_acesso(ws_processado[f"{COL_RESULTADO_AD}{row}"].value)
    acesso_ms = resultado_mensal_acesso(ws_processado[f"{COL_RESULTADO_MS}{row}"].value)
    return bool(acesso_ad and acesso_ms)


MENSAGEM_DADOS_INTERNOS_AUSENTES = "Não foram encontrados dados em sistemas internos"
SITUACAO_LEGAL_NAO_INFORMADA = "Situação legal não informada"
SITUACAO_LEGAL_CONSULTA_INDISPONIVEL = "Consulta indisponível"
SITUACAO_LEGAL_RESPOSTA_INVALIDA = "Resposta inválida do sistema interno"


def sanitizar_excel_texto(valor):
    if valor is None:
        return ""
    texto = str(valor)
    if texto and texto[0] in ("=", "+", "-", "@"):
        return "'" + texto
    return texto


def escrever_texto_excel(cell, valor, sanitizar_formula=True):
    if valor is None:
        cell.value = ""
    else:
        texto = str(valor)
        cell.value = sanitizar_excel_texto(texto) if sanitizar_formula else texto
    cell.data_type = "s"


def senha_para_chave(valor):
    if valor is None:
        return ""
    texto = str(valor)
    if len(texto) > 1 and texto[0] == "'" and texto[1] in ("=", "+", "-", "@"):
        return texto[1:]
    return texto


def cpf_para_chave_credencial(valor):
    cpf = normalizar_cpf_para_comparacao(valor)
    if cpf:
        return cpf
    digitos = re.sub(r"\D", "", str(valor or ""))
    return digitos if digitos else ""


def chave_credencial_mensal(cpf, senha):
    return (cpf_para_chave_credencial(cpf), senha_para_chave(senha))


def valor_ausente_mensal(valor):
    texto = "" if valor is None else str(valor).strip()
    if not texto:
        return True
    texto_normalizado = texto_sem_acentos(texto).lower()
    genericos = (
        "nome nao encontrado",
        "e-mail funcional nao encontrado",
        "email funcional nao encontrado",
        "cpf nao encontrado",
        "nao foram encontrados dados em sistemas internos",
        "nao encontrado",
        "none",
        "nan",
    )
    return any(item in texto_normalizado for item in genericos)


def valor_util_mensal(valor):
    texto = str(valor).strip()
    if valor_ausente_mensal(valor):
        return ""
    return texto


def preencher_ausente_mensal(valor):
    texto = valor_util_mensal(valor)
    return texto if texto else MENSAGEM_DADOS_INTERNOS_AUSENTES


def normalizar_url_vazamento(valor):
    texto = valor_util_mensal(valor)
    if not texto:
        return ""
    parsed = urlparse(texto)
    if parsed.scheme and parsed.netloc:
        return parsed._replace(query="", fragment="").geturl()
    return re.split(r"[?#]", texto, maxsplit=1)[0].strip()


def mes_numero_por_abreviacao(abreviacao):
    alvo = texto_sem_acentos(str(abreviacao or "")).upper()
    for numero, mes in MESES_PT.items():
        if mes == alvo:
            return numero
    return 99


def chave_data_identificacao(valor):
    match = re.fullmatch(r"(\d{2})([A-Z]{3})", str(valor or "").strip().upper())
    if not match:
        return (99, 99, str(valor or ""))
    return (mes_numero_por_abreviacao(match.group(2)), int(match.group(1)), "")


def combinar_datas_identificacao(datas_existentes, nova_data):
    datas = []
    for item in str(datas_existentes or "").split(","):
        data = item.strip().upper()
        if data and data not in datas:
            datas.append(data)
    nova = str(nova_data or "").strip().upper()
    if nova and nova not in datas:
        datas.append(nova)
    datas.sort(key=chave_data_identificacao)
    return ", ".join(datas)


def contar_datas_identificacao(datas):
    if not datas:
        return 0
    return len([item for item in str(datas).split(",") if item.strip()])


def capitalizar_legivel(valor):
    texto = re.sub(r"\s+", " ", str(valor or "").strip())
    if not texto:
        return ""
    return " ".join(parte[:1].upper() + parte[1:].lower() for parte in texto.split(" "))


def situacao_legal_valida(valor):
    texto = str(valor or "").strip()
    if not texto or texto in (
        MENSAGEM_DADOS_INTERNOS_AUSENTES,
        SITUACAO_LEGAL_NAO_INFORMADA,
        SITUACAO_LEGAL_CONSULTA_INDISPONIVEL,
        SITUACAO_LEGAL_RESPOSTA_INVALIDA,
    ):
        return False
    return bool(re.fullmatch(r".+\s-\s.+", texto))


def situacao_legal_mais_confiavel(valor_atual, valor_novo):
    novo = str(valor_novo or "").strip()
    atual = str(valor_atual or "").strip()
    if not novo:
        return atual
    if situacao_legal_valida(novo):
        return novo
    if situacao_legal_valida(atual):
        return atual
    if atual and atual not in (SITUACAO_LEGAL_CONSULTA_INDISPONIVEL, SITUACAO_LEGAL_RESPOSTA_INVALIDA):
        return atual
    return novo


def extrair_texto_situacao_legal(registro):
    situacao = registro.get("situacaoLegal") if isinstance(registro, dict) else None
    if not isinstance(situacao, dict):
        return ""
    descricao = capitalizar_legivel(situacao.get("descricao"))
    tipo = situacao.get("tipoSituacaoLegal")
    descricao_tipo = capitalizar_legivel(tipo.get("descricaoTipoSituacaoLegal") if isinstance(tipo, dict) else "")
    if descricao and descricao_tipo:
        return f"{descricao} - {descricao_tipo}"
    return ""


def selecionar_registro_situacao_legal(dados, cpf=None, linha=None):
    if len(dados) > 1:
        registrar_evento_auditoria(
            "identidade",
            linha=linha,
            sistema="API",
            cpf=cpf,
            operacao="situacao_legal",
            status="MULTIPLOS_REGISTROS",
            mensagem=f"API de situacao legal retornou {len(dados)} registros; aplicando criterio de selecao.",
        )
    for registro in dados:
        funcao = registro.get("funcao") if isinstance(registro, dict) else None
        if isinstance(funcao, dict) and funcao.get("principal") is True and str(funcao.get("status") or "").strip().upper() == "ATIVO":
            if extrair_texto_situacao_legal(registro):
                return registro
    for registro in dados:
        if extrair_texto_situacao_legal(registro):
            return registro
    return dados[0] if dados else None


def interpretar_resposta_situacao_legal(payload, cpf=None, linha=None):
    if not isinstance(payload, dict):
        return ResultadoOperacao(ERRO_INTERNO, SITUACAO_LEGAL_RESPOSTA_INVALIDA, SITUACAO_LEGAL_RESPOSTA_INVALIDA)
    if payload.get("executouComSucesso") is not True or payload.get("executouComFalha") is not False:
        return ResultadoOperacao(ERRO_INTERNO, SITUACAO_LEGAL_RESPOSTA_INVALIDA, SITUACAO_LEGAL_RESPOSTA_INVALIDA)
    dados = payload.get("dados")
    if not isinstance(dados, list):
        return ResultadoOperacao(ERRO_INTERNO, SITUACAO_LEGAL_RESPOSTA_INVALIDA, SITUACAO_LEGAL_RESPOSTA_INVALIDA)
    if not dados:
        return ResultadoOperacao(NAO_ENCONTRADO, MENSAGEM_DADOS_INTERNOS_AUSENTES, MENSAGEM_DADOS_INTERNOS_AUSENTES)
    registro = selecionar_registro_situacao_legal(dados, cpf=cpf, linha=linha)
    texto = extrair_texto_situacao_legal(registro)
    if not texto:
        return ResultadoOperacao(INCONCLUSIVO, SITUACAO_LEGAL_NAO_INFORMADA, SITUACAO_LEGAL_NAO_INFORMADA)
    return ResultadoOperacao(SUCESSO, "Situação legal encontrada.", texto)


def consultar_situacao_legal_por_cpf(cpf, session=None, linha=None):
    cpf_normalizado = normalizar_cpf_para_comparacao(cpf)
    if not cpf_normalizado:
        registrar_evento_auditoria("identidade", linha=linha, sistema="API", cpf=cpf, operacao="situacao_legal", status=ERRO_DE_ENTRADA, mensagem="CPF invalido para consulta de situacao legal.")
        return ResultadoOperacao(ERRO_DE_ENTRADA, "CPF inválido.", "")
    sessao = session or requests.Session()
    url = f"https://siteapi.com.br/api/v1/PolicialMilitar/cpf/{cpf_normalizado}/informacao"
    temporarios = {502, 503, 504}
        try:
            for tentativa in range(1, 4):
            try:
                response = sessao.get(url, verify=False, timeout=(5, 15))
                if response.status_code in temporarios:
                    if tentativa < 3:
                        time.sleep(0.5 * tentativa)
                        continue
                    registrar_evento_auditoria("identidade", linha=linha, sistema="API", cpf=cpf_normalizado, operacao="situacao_legal", status=SISTEMA_INDISPONIVEL, mensagem=f"HTTP temporario {response.status_code} apos tentativas permitidas.")
                    return ResultadoOperacao(SISTEMA_INDISPONIVEL, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL)
                if response.status_code < 200 or response.status_code >= 300:
                    registrar_evento_auditoria("identidade", linha=linha, sistema="API", cpf=cpf_normalizado, operacao="situacao_legal", status=ERRO_DE_REDE, mensagem=f"HTTP inesperado {response.status_code} na consulta de situacao legal.")
                    return ResultadoOperacao(ERRO_DE_REDE, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL if response.status_code >= 500 else SITUACAO_LEGAL_RESPOSTA_INVALIDA, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL if response.status_code >= 500 else SITUACAO_LEGAL_RESPOSTA_INVALIDA)
                try:
                    payload = response.json()
                except ValueError:
                    registrar_evento_auditoria("identidade", linha=linha, sistema="API", cpf=cpf_normalizado, operacao="situacao_legal", status=ERRO_INTERNO, mensagem="JSON invalido na consulta de situacao legal.")
                    return ResultadoOperacao(ERRO_INTERNO, SITUACAO_LEGAL_RESPOSTA_INVALIDA, SITUACAO_LEGAL_RESPOSTA_INVALIDA)
                resultado = interpretar_resposta_situacao_legal(payload, cpf=cpf_normalizado, linha=linha)
                if resultado.status != SUCESSO:
                    registrar_evento_auditoria("identidade", linha=linha, sistema="API", cpf=cpf_normalizado, operacao="situacao_legal", status=resultado.status, mensagem=resultado.mensagem)
                return resultado
            except requests.Timeout:
                if tentativa < 3:
                    time.sleep(0.5 * tentativa)
                    continue
                registrar_evento_auditoria("identidade", linha=linha, sistema="API", cpf=cpf_normalizado, operacao="situacao_legal", status=TIMEOUT, mensagem="Timeout apos tentativas permitidas na consulta de situacao legal.")
                return ResultadoOperacao(TIMEOUT, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL)
            except requests.ConnectionError:
                if tentativa < 3:
                    time.sleep(0.5 * tentativa)
                    continue
                registrar_evento_auditoria("identidade", linha=linha, sistema="API", cpf=cpf_normalizado, operacao="situacao_legal", status=ERRO_DE_REDE, mensagem="Falha de conexao apos tentativas permitidas na consulta de situacao legal.")
                return ResultadoOperacao(ERRO_DE_REDE, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL)
            except requests.RequestException as exc:
                registrar_evento_auditoria("identidade", linha=linha, sistema="API", cpf=cpf_normalizado, operacao="situacao_legal", status=ERRO_DE_REDE, mensagem=f"Falha HTTP sanitizada na consulta de situacao legal: {higienizar_mensagem_log(exc)}")
                return ResultadoOperacao(ERRO_DE_REDE, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL)
        return ResultadoOperacao(ERRO_INTERNO, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL, SITUACAO_LEGAL_CONSULTA_INDISPONIVEL)
    finally:
        if session is None:
            sessao.close()


def aplicar_cache_situacao_legal(registros, session=None):
    cache = {}
    for registro in registros:
        cpf = normalizar_cpf_para_comparacao(registro.get("cpf"))
        if not cpf:
            registro["situacao_legal"] = ""
            continue
        if cpf not in cache:
            cache[cpf] = consultar_situacao_legal_por_cpf(cpf, session=session, linha=registro.get("linha")).valor
        registro["situacao_legal"] = cache[cpf]
    return cache


def extrair_registros_mensais_unificados(ws_processado, sistema=None):
    registros = []
    for row in range(2, ws_processado.max_row + 1):
        if resultados_nt_em_ambos(ws_processado, row):
            continue
        if not registro_mensal_processavel(ws_processado, row):
            continue
        registros.append({
            "linha": row,
            "nome": preencher_ausente_mensal(ws_processado[f"{COL_NOME}{row}"].value),
            "cpf": normalizar_cpf_para_comparacao(ws_processado[f"{COL_CPF_RESOLVIDO}{row}"].value) or ws_processado[f"{COL_CPF_RESOLVIDO}{row}"].value,
            "senha": senha_para_chave(ws_processado[f"{COL_SENHA}{row}"].value),
            "email": preencher_ausente_mensal(ws_processado[f"{COL_EMAIL_FUNCIONAL}{row}"].value),
            "url": normalizar_url_vazamento(ws_processado[f"{COL_URL_VAZAMENTO}{row}"].value),
            "acesso_ad": resultado_mensal_acesso(ws_processado[f"{COL_RESULTADO_AD}{row}"].value),
            "acesso_ms": resultado_mensal_acesso(ws_processado[f"{COL_RESULTADO_MS}{row}"].value),
            "situacao_legal": "",
        })
    return registros


def configurar_aba_mensal_unificada(ws, headers):
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col)
        cell.value = header
        cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"
    # O filtro fica dentro da Tabela oficial do Excel. Manter AutoFilter
    # tambem na planilha gera reparo/remocao de tabela ao abrir no Excel.
    ws.auto_filter.ref = None


def limpar_tabelas_excel(ws):
    try:
        for nome in list(ws.tables.keys()):
            del ws.tables[nome]
    except AttributeError:
        ws._tables = []


def aplicar_tabela_excel(ws, nome_tabela, headers):
    limpar_tabelas_excel(ws)
    ws.auto_filter.ref = None
    ultima_linha = max(ws.max_row, 1)
    ultima_coluna = chr(64 + len(headers))
    ref = f"A1:{ultima_coluna}{ultima_linha}"
    tabela = Table(displayName=nome_tabela, ref=ref)
    estilo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tabela.tableStyleInfo = estilo
    ws.add_table(tabela)


def aplicar_formatacao_mensal_unificada(wb):
    verde = PatternFill("solid", fgColor="C6EFCE")
    vermelho = PatternFill("solid", fgColor="FFC7CE")
    widths_cred = {"A": 34, "B": 16, "C": 28, "D": 38, "E": 58, "F": 16, "G": 28, "H": 38, "I": 16, "J": 14, "K": 14, "L": 28, "M": 34, "N": 18, "O": 34}
    widths_total = {"A": 34, "B": 16, "C": 28, "D": 38, "E": 58, "F": 16, "G": 28, "H": 38, "I": 16, "J": 14, "K": 14, "L": 28}
    for nome_aba in (ABAS_MENSAIS["AD"], ABAS_MENSAIS["MS"]):
        ws = wb[nome_aba]
        configurar_aba_mensal_unificada(ws, HEADERS_MENSAL_CREDENCIAIS)
        for coluna, largura in widths_cred.items():
            ws.column_dimensions[coluna].width = largura
        for row in range(2, ws.max_row + 1):
            for col in range(1, len(HEADERS_MENSAL_CREDENCIAIS) + 1):
                ws.cell(row, col).font = Font(name="Arial", size=10)
                ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=col in (1, 4, 5, 8, 12, 13, 15))
            for col in (10, 11):
                cell = ws.cell(row, col)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(name="Arial", size=10, bold=True)
                if cell.value == RESULTADO_SIM:
                    cell.fill = verde
                elif cell.value == RESULTADO_NAO:
                    cell.fill = vermelho
        aplicar_tabela_excel(ws, "Tabela_" + re.sub(r"\W+", "", nome_aba), HEADERS_MENSAL_CREDENCIAIS)
    ws_total = wb[ABAS_MENSAIS["Total"]]
    configurar_aba_mensal_unificada(ws_total, HEADERS_MENSAL_TOTAL)
    for coluna, largura in widths_total.items():
        ws_total.column_dimensions[coluna].width = largura
    for row in range(2, ws_total.max_row + 1):
        for col in range(1, len(HEADERS_MENSAL_TOTAL) + 1):
            ws_total.cell(row, col).font = Font(name="Arial", size=10)
            ws_total.cell(row, col).alignment = Alignment(vertical="top", wrap_text=col in (1, 4, 5, 8, 12))
        for col in (10, 11):
            cell = ws_total.cell(row, col)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Arial", size=10, bold=True)
            if cell.value == RESULTADO_SIM:
                cell.fill = verde
            elif cell.value == RESULTADO_NAO:
                cell.fill = vermelho
    aplicar_tabela_excel(ws_total, "Tabela_Total", HEADERS_MENSAL_TOTAL)


def garantir_abas_mensais_unificadas(wb):
    criadas = []
    aliases = {"Credenciais_AD": ABAS_MENSAIS["AD"], "Credenciais_MS": ABAS_MENSAIS["MS"]}
    for antigo, novo in aliases.items():
        if antigo in wb.sheetnames and novo not in wb.sheetnames:
            wb[antigo].title = novo
    for sistema in ("AD", "MS"):
        nome = ABAS_MENSAIS[sistema]
        if nome not in wb.sheetnames:
            wb.create_sheet(nome)
            criadas.append(nome)
        migrar_aba_para_layout_mensal(wb[nome], HEADERS_MENSAL_CREDENCIAIS)
    if ABAS_MENSAIS["Total"] not in wb.sheetnames:
        wb.create_sheet(ABAS_MENSAIS["Total"])
        criadas.append(ABAS_MENSAIS["Total"])
    migrar_aba_para_layout_mensal(wb[ABAS_MENSAIS["Total"]], HEADERS_MENSAL_TOTAL)
    for nome in list(wb.sheetnames):
        ws = wb[nome]
        if nome not in ABAS_MENSAIS.values() and len(wb.sheetnames) > 3 and ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
            del wb[nome]
    # A exigencia atual e exatamente tres abas mensais; abas antigas equivalentes ja foram renomeadas.
    for nome in list(wb.sheetnames):
        if nome not in ABAS_MENSAIS.values():
            del wb[nome]
    aplicar_formatacao_mensal_unificada(wb)
    return criadas


def criar_workbook_mensal_unificado():
    wb = Workbook()
    wb.active.title = ABAS_MENSAIS["AD"]
    wb.create_sheet(ABAS_MENSAIS["MS"])
    wb.create_sheet(ABAS_MENSAIS["Total"])
    garantir_abas_mensais_unificadas(wb)
    return wb


def mapa_headers(ws):
    return {normalizar_header(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1) if ws.cell(1, col).value}


def ler_por_header(ws, row, header, mapa):
    col = mapa.get(normalizar_header(header))
    return ws.cell(row, col).value if col else None


def migrar_aba_para_layout_mensal(ws, headers_destino):
    if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
        configurar_aba_mensal_unificada(ws, headers_destino)
        return
    mapa = mapa_headers(ws)
    linhas = []
    for row in range(2, ws.max_row + 1):
        if not any(ws.cell(row, col).value is not None for col in range(1, ws.max_column + 1)):
            continue
        senha = ler_por_header(ws, row, "SENHA", mapa) or ler_por_header(ws, row, "Senha", mapa)
        registro = {
            "NOME": ler_por_header(ws, row, "NOME", mapa) or ler_por_header(ws, row, "Nome", mapa),
            "CPF": normalizar_cpf_para_comparacao(ler_por_header(ws, row, "CPF", mapa)) or ler_por_header(ws, row, "CPF", mapa),
            "SENHA": senha_para_chave(senha),
            "EMAIL": ler_por_header(ws, row, "EMAIL", mapa) or ler_por_header(ws, row, "E-mail", mapa),
            "URL": ler_por_header(ws, row, "URL", mapa) or ler_por_header(ws, row, "URL de vazamento", mapa),
            "DATA COLETA": ler_por_header(ws, row, "DATA COLETA", mapa),
            "QUANTIDADE DE IDENTIFICAÇÕES": ler_por_header(ws, row, "QUANTIDADE DE IDENTIFICAÇÕES", mapa) or ler_por_header(ws, row, "QUANTIDADE DE IDENTIFICA??ES", mapa) or ler_por_header(ws, row, HEADER_CONTAGEM_IDENTIFICACOES, mapa),
            "DATA DAS IDENTIFICAÇÕES": ler_por_header(ws, row, "DATA DAS IDENTIFICAÇÕES", mapa) or ler_por_header(ws, row, "DATA DAS IDENTIFICA??ES", mapa) or ler_por_header(ws, row, "Datas das identificações", mapa) or ler_por_header(ws, row, "Datas das identifica??es", mapa),
            "FONTE": ler_por_header(ws, row, "FONTE", mapa) or "BTT-APURA",
            "ACESSO AD": ler_por_header(ws, row, "ACESSO AD", mapa) or ler_por_header(ws, row, "AD", mapa),
            "ACESSO MS": ler_por_header(ws, row, "ACESSO MS", mapa) or ler_por_header(ws, row, "MS", mapa),
            "SITUAÇÃO LEGAL": ler_por_header(ws, row, "SITUAÇÃO LEGAL", mapa) or ler_por_header(ws, row, "SITUACAO LEGAL", mapa),
            "OBSERVAÇÕES": ler_por_header(ws, row, "OBSERVAÇÕES", mapa) or ler_por_header(ws, row, "OBSERVA??ES", mapa),
            "RDS": ler_por_header(ws, row, "RDS", mapa),
            "MSG BLOQUEIO": ler_por_header(ws, row, "MSG BLOQUEIO", mapa),
        }
        datas = registro["DATA DAS IDENTIFICAÇÕES"]
        if not datas:
            data_origem = ler_por_header(ws, row, "Data desta identificação", mapa) or ler_por_header(ws, row, "Data desta identifica??o", mapa) or ler_por_header(ws, row, "Primeira identificação", mapa) or ler_por_header(ws, row, "Primeira identifica??o", mapa)
            if isinstance(data_origem, datetime):
                datas = formatar_data_identificacao(data_origem)
        registro["DATA DAS IDENTIFICAÇÕES"] = combinar_datas_identificacao("", datas) if datas else ""
        if not registro["QUANTIDADE DE IDENTIFICAÇÕES"]:
            registro["QUANTIDADE DE IDENTIFICAÇÕES"] = contar_datas_identificacao(registro["DATA DAS IDENTIFICAÇÕES"])
        linhas.append(registro)
    ws.delete_rows(1, ws.max_row)
    configurar_aba_mensal_unificada(ws, headers_destino)
    vistos = set()
    for registro in linhas:
        chave = chave_credencial_mensal(registro.get("CPF"), registro.get("SENHA"))
        if chave in vistos:
            continue
        vistos.add(chave)
        row = ws.max_row + 1
        for col, header in enumerate(headers_destino, start=1):
            escrever_texto_excel(ws.cell(row, col), registro.get(header, ""), sanitizar_formula=(header != "SENHA"))
        datas = ws.cell(row, 8).value
        ws.cell(row, 7).value = contar_datas_identificacao(datas)


def indexar_aba_por_credencial(ws):
    indice = {}
    for row in range(2, ws.max_row + 1):
        chave = chave_credencial_mensal(ws.cell(row, 2).value, ws.cell(row, 3).value)
        if chave not in indice:
            indice[chave] = row
    return indice


def preencher_base_credencial(ws, row, registro, data_label, incrementar_data):
    if valor_ausente_mensal(ws.cell(row, 1).value) and valor_util_mensal(registro.get("nome")):
        escrever_texto_excel(ws.cell(row, 1), valor_util_mensal(registro.get("nome")))
    elif not valor_util_mensal(ws.cell(row, 1).value):
        escrever_texto_excel(ws.cell(row, 1), MENSAGEM_DADOS_INTERNOS_AUSENTES)
    escrever_texto_excel(ws.cell(row, 2), normalizar_cpf_para_comparacao(registro.get("cpf")) or registro.get("cpf"))
    escrever_texto_excel(ws.cell(row, 3), registro.get("senha"), sanitizar_formula=False)
    if valor_ausente_mensal(ws.cell(row, 4).value) and valor_util_mensal(registro.get("email")):
        escrever_texto_excel(ws.cell(row, 4), valor_util_mensal(registro.get("email")))
    elif not valor_util_mensal(ws.cell(row, 4).value):
        escrever_texto_excel(ws.cell(row, 4), MENSAGEM_DADOS_INTERNOS_AUSENTES)
    if not valor_util_mensal(ws.cell(row, 5).value):
        escrever_texto_excel(ws.cell(row, 5), normalizar_url_vazamento(registro.get("url")))
    if not valor_util_mensal(ws.cell(row, 6).value):
        escrever_texto_excel(ws.cell(row, 6), data_label)
    datas_antes = ws.cell(row, 8).value
    datas_depois = combinar_datas_identificacao(datas_antes, data_label if incrementar_data else "")
    escrever_texto_excel(ws.cell(row, 8), datas_depois)
    ws.cell(row, 7).value = contar_datas_identificacao(datas_depois)
    escrever_texto_excel(ws.cell(row, 9), "BTT-APURA")
    escrever_texto_excel(ws.cell(row, 10), registro.get("acesso_ad"))
    escrever_texto_excel(ws.cell(row, 11), registro.get("acesso_ms"))
    escrever_texto_excel(ws.cell(row, 12), situacao_legal_mais_confiavel(ws.cell(row, 12).value, registro.get("situacao_legal")))


def upsert_linha_mensal(ws, registro, data_label, estatisticas, controle_dia, incluir_manuais):
    indice = indexar_aba_por_credencial(ws)
    chave = chave_credencial_mensal(registro.get("cpf"), registro.get("senha"))
    row = indice.get(chave)
    if row is None:
        row = ws.max_row + 1
        estatisticas["novos"] += 1
        incrementar_data = chave not in controle_dia
    else:
        estatisticas["atualizados"] += 1
        datas_existentes = str(ws.cell(row, 8).value or "")
        incrementar_data = data_label not in [d.strip().upper() for d in datas_existentes.split(",")]
    if incrementar_data:
        controle_dia.add(chave)
        estatisticas["datas_acrescentadas"] += 1
    preencher_base_credencial(ws, row, registro, data_label, incrementar_data)
    if incluir_manuais:
        # As colunas 13, 14 e 15 sao manuais e nunca sao sobrescritas.
        for col in (13, 14, 15):
            if ws.cell(row, col).value is None:
                ws.cell(row, col).value = ""
    estatisticas["contadores_incrementados"] += 1 if incrementar_data else 0
    return row


def sincronizar_abas_positivas(wb):
    total = wb[ABAS_MENSAIS["Total"]]
    for sistema, aba in (("AD", ABAS_MENSAIS["AD"]), ("MS", ABAS_MENSAIS["MS"])):
        ws = wb[aba]
        manuais_por_chave = {}
        for row in range(2, ws.max_row + 1):
            chave_existente = chave_credencial_mensal(ws.cell(row, 2).value, ws.cell(row, 3).value)
            if chave_existente[0] or chave_existente[1]:
                manuais_por_chave[chave_existente] = {
                    13: ws.cell(row, 13).value,
                    14: ws.cell(row, 14).value,
                    15: ws.cell(row, 15).value,
                }
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        col_acesso = 10 if sistema == "AD" else 11
        for row in range(2, total.max_row + 1):
            if total.cell(row, col_acesso).value != RESULTADO_SIM:
                continue
            chave = chave_credencial_mensal(total.cell(row, 2).value, total.cell(row, 3).value)
            destino = ws.max_row + 1
            for col in range(1, 13):
                ws.cell(destino, col).value = total.cell(row, col).value
                ws.cell(destino, col).data_type = total.cell(row, col).data_type
            manuais = manuais_por_chave.get(chave, {})
            for col in (13, 14, 15):
                ws.cell(destino, col).value = manuais.get(col) or ""


def gerar_total_a_partir_abas_positivas_migradas(wb):
    total = wb[ABAS_MENSAIS["Total"]]
    if total.max_row > 1:
        return 0
    indice_total = {}
    importados = 0
    for sistema, aba in (("AD", ABAS_MENSAIS["AD"]), ("MS", ABAS_MENSAIS["MS"])):
        ws = wb[aba]
        col_acesso = 10 if sistema == "AD" else 11
        outro_col_acesso = 11 if sistema == "AD" else 10
        for row in range(2, ws.max_row + 1):
            chave = chave_credencial_mensal(ws.cell(row, 2).value, ws.cell(row, 3).value)
            if not (chave[0] or chave[1]):
                continue
            destino = indice_total.get(chave)
            if destino is None:
                destino = total.max_row + 1
                indice_total[chave] = destino
                importados += 1
                for col in range(1, 13):
                    total.cell(destino, col).value = ws.cell(row, col).value
                    total.cell(destino, col).data_type = ws.cell(row, col).data_type
            else:
                datas = combinar_datas_identificacao(total.cell(destino, 8).value, ws.cell(row, 8).value)
                escrever_texto_excel(total.cell(destino, 8), datas)
                total.cell(destino, 7).value = contar_datas_identificacao(datas)
                for col in (1, 4, 5, 6):
                    if not valor_util_mensal(total.cell(destino, col).value) and valor_util_mensal(ws.cell(row, col).value):
                        total.cell(destino, col).value = ws.cell(row, col).value
                        total.cell(destino, col).data_type = ws.cell(row, col).data_type
            if total.cell(destino, col_acesso).value not in (RESULTADO_SIM, RESULTADO_NAO):
                escrever_texto_excel(total.cell(destino, col_acesso), RESULTADO_SIM)
            if total.cell(destino, outro_col_acesso).value not in (RESULTADO_SIM, RESULTADO_NAO):
                escrever_texto_excel(total.cell(destino, outro_col_acesso), RESULTADO_NAO)
    return importados


def salvar_workbook_mensal_unificado_atomicamente(wb, destino):
    diretorio = os.path.dirname(destino)
    temp_path = None
    try:
        fd, temp_path = tempfile.mkstemp(prefix=".mensal_", suffix=".xlsx", dir=diretorio)
        os.close(fd)
        aplicar_formatacao_mensal_unificada(wb)
        wb.save(temp_path)
        teste = load_workbook(temp_path, read_only=True)
        for nome in ABAS_MENSAIS.values():
            if nome not in teste.sheetnames:
                teste.close()
                raise RuntimeError(f"Aba mensal obrigatoria ausente: {nome}")
        if [teste[ABAS_MENSAIS["AD"]].cell(1, col).value for col in range(1, len(HEADERS_MENSAL_CREDENCIAIS) + 1)] != HEADERS_MENSAL_CREDENCIAIS:
            teste.close()
            raise RuntimeError("Cabecalhos da aba Credenciais AD ficaram invalidos.")
        if [teste[ABAS_MENSAIS["MS"]].cell(1, col).value for col in range(1, len(HEADERS_MENSAL_CREDENCIAIS) + 1)] != HEADERS_MENSAL_CREDENCIAIS:
            teste.close()
            raise RuntimeError("Cabecalhos da aba Credenciais MS ficaram invalidos.")
        if [teste[ABAS_MENSAIS["Total"]].cell(1, col).value for col in range(1, len(HEADERS_MENSAL_TOTAL) + 1)] != HEADERS_MENSAL_TOTAL:
            teste.close()
            raise RuntimeError("Cabecalhos da aba Total ficaram invalidos.")
        teste.close()
        os.replace(temp_path, destino)
    except PermissionError:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass
        raise PermissionError(f"Nao foi possivel atualizar {os.path.basename(destino)} porque o arquivo esta aberto ou sem permissao de escrita.")
    except Exception:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass
        raise


def caminho_arquivo_mensal_unificado(data_execucao):
    return os.path.join(script_dir, f"Credenciais_{obter_competencia_atual(data_execucao)}.xlsx")


def carregar_ou_criar_arquivo_mensal_unificado(caminho):
    criado = not os.path.exists(caminho)
    wb = criar_workbook_mensal_unificado() if criado else load_workbook(caminho)
    abas_criadas = garantir_abas_mensais_unificadas(wb)
    gerar_total_a_partir_abas_positivas_migradas(wb)
    return wb, criado, abas_criadas


def migrar_arquivos_mensais_antigos(wb, data_execucao, estatisticas):
    # Arquivos antigos separados sao preservados; quando o mensal novo nao existe,
    # dados legados sao lidos pelas abas renomeadas/normalizadas se presentes.
    return []


def nova_estatistica_mensal():
    return {
        "novos": 0,
        "atualizados": 0,
        "contadores_incrementados": 0,
        "datas_acrescentadas": 0,
        "resultados_atualizados": 0,
        "urls_acrescentadas": 0,
    }


def atualizar_situacao_legal_existente_total(wb, cache, session=None):
    total = wb[ABAS_MENSAIS["Total"]]
    atualizados = 0
    for row in range(2, total.max_row + 1):
        cpf = normalizar_cpf_para_comparacao(total.cell(row, 2).value)
        if not cpf:
            continue
        atual = total.cell(row, 12).value
        if situacao_legal_valida(atual):
            continue
        if cpf not in cache:
            cache[cpf] = consultar_situacao_legal_por_cpf(cpf, session=session, linha="-").valor
        novo = situacao_legal_mais_confiavel(atual, cache[cpf])
        if novo != atual:
            escrever_texto_excel(total.cell(row, 12), novo)
            atualizados += 1
    return atualizados


def atualizar_consolidacoes_mensais(ws_processado, data_execucao, session_situacao_legal=None):
    caminho = caminho_arquivo_mensal_unificado(data_execucao)
    estatisticas = {"AD": nova_estatistica_mensal(), "MS": nova_estatistica_mensal(), "Total": nova_estatistica_mensal()}
    registros = extrair_registros_mensais_unificados(ws_processado)
    registros_ad_sim = [r for r in registros if r["acesso_ad"] == RESULTADO_SIM]
    registros_ms_sim = [r for r in registros if r["acesso_ms"] == RESULTADO_SIM]
    resultado = {
        "arquivo": caminho,
        "situacao": "",
        "abas_criadas": [],
        "registros_recebidos": len(registros),
        "positivos_ad": len(registros_ad_sim),
        "positivos_ms": len(registros_ms_sim),
        "total_enviados": len(registros),
        "migracao": [],
        "erro": "",
    }
    try:
        wb, criado, abas_criadas = carregar_ou_criar_arquivo_mensal_unificado(caminho)
        resultado["abas_criadas"] = abas_criadas
        resultado["situacao"] = "arquivo criado" if criado else "arquivo existente atualizado"
        cache_situacao_legal = aplicar_cache_situacao_legal(registros, session=session_situacao_legal)
        atualizar_situacao_legal_existente_total(wb, cache_situacao_legal, session=session_situacao_legal)
        data_label = formatar_data_identificacao(data_execucao)
        controle_total = set()
        controle_ad = set()
        controle_ms = set()
        for registro in registros:
            upsert_linha_mensal(wb[ABAS_MENSAIS["Total"]], registro, data_label, estatisticas["Total"], controle_total, False)
            if registro["acesso_ad"] == RESULTADO_SIM:
                upsert_linha_mensal(wb[ABAS_MENSAIS["AD"]], registro, data_label, estatisticas["AD"], controle_ad, True)
            if registro["acesso_ms"] == RESULTADO_SIM:
                upsert_linha_mensal(wb[ABAS_MENSAIS["MS"]], registro, data_label, estatisticas["MS"], controle_ms, True)
        sincronizar_abas_positivas(wb)
        salvar_workbook_mensal_unificado_atomicamente(wb, caminho)
        wb.close()
    except Exception as exc:
        resultado["erro"] = higienizar_mensagem_log(str(exc))
    for sistema in ("AD", "MS"):
        encontrados = len(registros_ad_sim) if sistema == "AD" else len(registros_ms_sim)
        resultado[sistema] = {
            "sistema": sistema,
            "arquivo": caminho,
            "encontrados": encontrados,
            "adicionados": estatisticas[sistema]["novos"],
            "duplicados": estatisticas[sistema]["atualizados"],
            "erro": resultado["erro"],
            **estatisticas[sistema],
        }
    resultado["Total"] = {"arquivo": caminho, "encontrados": len(registros), **estatisticas["Total"]}
    return resultado

def main():
    execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
    inicio = datetime.now()
    iniciar_contexto_logs(execution_id, inicio)
    logger_identidade = configurar_logger("buscar_nome", os.path.join(script_dir, "Logs_execucao.txt"), execution_id)
    logger_auth = configurar_logger("testar_credenciais", os.path.join(script_dir, "Logs_auditoria.txt"), execution_id)
    resumo = Counter()
    resumo["execution_id"] = execution_id
    resumo["arquivo_entrada"] = arquivo_excel
    resumo["inicio"] = inicio
    cliente_api = None
    try:
        wb, ws, layout_novo = validar_arquivo_entrada(arquivo_excel)
        preparar_layout(ws, layout_novo)
        resumo["total_linhas"] = ws.max_row - 1
        cliente_api = ClienteAPIsInternas(logger_identidade)
        processar_identidade(ws, cliente_api, logger_identidade, resumo)
        processar_autenticacoes(ws, logger_auth, resumo)
        nome_final = f"Credenciais_{datetime.now().strftime('%d-%m-%Y_%H-%M')}.xlsx"
        resultado_final = os.path.join(script_dir, nome_final)
        resumo["arquivo_final"] = resultado_final
        resumo["termino"] = datetime.now()
        wb_saida = criar_workbook_saida(ws)
        salvar_atomicamente(wb_saida, resultado_final)
        consolidacao_mensal = atualizar_consolidacoes_mensais(ws, resumo["termino"])
        consolidacao_mensal["auditoria_nt"] = {
            "total_linhas": resumo.get("total_linhas", 0),
            "linhas_removidas": resumo.get("linhas_removidas_nt_numeros", []),
            "linhas_mantidas": resumo.get("linhas_mantidas_pos_filtro", max(ws.max_row - 1, 0)),
        }
        escrever_logs_finais(resumo, resultado_final, consolidacao_mensal)
        for sistema in ("AD", "MS"):
            info = consolidacao_mensal.get(sistema, {})
            print(f"\n{sistema}:")
            print(f"Resultados SIM encontrados nesta execução: {info.get('encontrados', 0)}")
            print(f"Novas linhas adicionadas: {info.get('adicionados', 0)}")
            print(f"Registros atualizados por identificacao repetida neste mes: {info.get('duplicados', 0)}")
            if info.get("erro"):
                print(f"Falha ao atualizar arquivo mensal: {info['erro']}")
            elif info.get("arquivo"):
                print(f"Arquivo mensal: {info['arquivo']}")
            else:
                print("Arquivo mensal: não criado nesta execução")
        print("\n📊 Resumo final")
        print(f"📁 Arquivo final: {resultado_final}")
        print(f"✅ Identidades resolvidas: {resumo['linhas_processadas']}")
        print(f"⚠️ Linhas ignoradas: {resumo['linhas_ignoradas']}")
        print(f"Linhas removidas por N/T em MS e AD: {resumo.get('linhas_removidas_nt', 0)}")
        print(f"🟢 MS sucesso: {resumo.get('MS_SUCESSO', 0)}")
        print(f"🟢 AD sucesso: {resumo.get('AD_SUCESSO', 0)}")
        print("🏁 Execução concluída.")
        return 0
    except Exception as exc:
        resumo["erro_critico"] = higienizar_mensagem_log(str(exc))
        resumo["termino"] = datetime.now()
        resumo.setdefault("arquivo_final", "")
        escrever_logs_finais(resumo, resumo.get("arquivo_final", ""), {"arquivo": "", "erro": str(exc)})
        print(f"Erro fatal: {exc}")
        return 1
    finally:
        if cliente_api:
            cliente_api.fechar()


if __name__ == "__main__":
    sys.exit(main())
